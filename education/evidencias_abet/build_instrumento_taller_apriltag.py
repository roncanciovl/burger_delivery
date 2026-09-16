#!/usr/bin/env python3
"""
Constructor para INSTRUMENTO_ABET_TALLER_LOCALIZACION_APRILTAG_MICROROS.docx
Genera el instrumento Word único de captura de evidencia, tablas de datos experimentales,
cuestionario de análisis técnico, evaluación y consolidación ABET para el Taller de
Localización Visual 2D con AprilTags (Kinova) y Lazo de Pose con micro-ROS en ESP32.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
TARGET = ROOT / "education" / "evidencias_abet" / "INSTRUMENTO_ABET_TALLER_LOCALIZACION_APRILTAG_MICROROS.docx"

FONT = "Arial"
BLUE = "1F4E79"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F2F2"

INDICATOR_11 = (
    "Indicador de desempeño 1.1: Formula y conecta el flujo de datos entre la cámara del Kinova, "
    "un nodo de localización en la PC y un nodo de control embebido en el microcontrolador."
)
INDICATOR_22 = (
    "Indicador de desempeño 2.2: Selecciona interfaces y mensajes compactos (imagen comprimida, "
    "geometry_msgs/msg/Pose2D) respetando las restricciones de ancho de banda y memoria del ESP32, "
    "y justifica sus limitaciones."
)
INDICATOR_61 = (
    "Indicador de desempeño 6.1: Diseña y ejecuta pruebas de lazo cerrado, midiendo la exactitud "
    "de la pose, su frecuencia de actualización y la respuesta del nodo embebido."
)
INDICATOR_64 = (
    "Indicador de desempeño 6.4: Interpreta fallas y diagnósticos experimentales aplicando protocolos "
    "de diagnóstico por capas para aislar errores en hardware y software."
)
INDICATOR_31_33 = (
    "Indicador de desempeño 3.1 - 3.3: Elabora documentación técnica reproducible del sistema ROS 2 "
    "y comunica resultados experimentales de percepción, calibración y control embebido."
)
INDICATOR_51 = (
    "Indicador de desempeño 5.1: Define roles técnicos, coordina la ejecución en equipo y fomenta "
    "un entorno de trabajo colaborativo e inclusivo."
)

LEVELS = [
    ("N5", "475–500", "Excelente: solución y evidencia completas, precisas, reproducibles y explicadas con profundidad analítica y matemática sobresaliente."),
    ("N4", "400–474", "Bueno: desempeño técnico correcto con omisiones menores que no comprometen la operación, reproducibilidad ni diagnóstico."),
    ("N3", "300–399", "Aceptable: demuestra el desempeño esencial con evidencia verificable. Es el umbral individual de logro."),
    ("N2", "150–299", "Cumplimiento parcial: evidencia incompleta, métricas faltantes o errores conceptuales en homografía o micro-ROS."),
    ("N1", "0–149", "No cumple: evidencia mínima, fragmentaria o no funcional. Sin evidencia obligatoria se registra 0."),
]

CRITERIA = [
    {
        "code": "C1",
        "name": "Arquitectura y lazo distribuido XRCE-DDS / DDS",
        "weight": 25,
        "so": "Student Outcome SO1",
        "indicator": INDICATOR_11,
        "descriptors": [
            "Además de N4, analiza la estructura de tramas XRCE-DDS, optimiza la memoria del microcontrolador ajustando el número de handles del executor y justifica arquitectónicamente la ausencia del agente en el grafo.",
            "Establece el lazo bidireccional completo entre la cámara, el localizador, el agente UDP y el ESP32, verificando Subscription count: 1 y documentando el flujo en rqt_graph.",
            "Configura el localizador y el agente micro-ROS, logrando que el ESP32 reciba /pose2d y responda con /distance_to_goal en Modo A o B.",
            "El lazo opera de forma intermitente, presenta problemas de direccionamiento IP o no logra que el ESP32 publique la distancia calculada.",
            "No logra establecer la comunicación entre la PC y el ESP32 o carece de evidencias funcionales obligatorias.",
        ],
    },
    {
        "code": "C2",
        "name": "Selección de interfaces compactas y homografía",
        "weight": 25,
        "so": "Student Outcome SO2",
        "indicator": INDICATOR_22,
        "descriptors": [
            "Además de N4, calcula analíticamente la matriz de homografía 3x3, demuestra la invarianza proyectiva ante cambios de perspectiva y propone el empaquetado de metadatos temporales mínimos en interfaces embebidas.",
            "Calibra con pie de rey el tamaño exacto del tag, orienta correctamente los ejes de tag_mesa, justifica el uso de Pose2D y calcula a mano la pose de prueba (Ejercicio 3.1).",
            "Configura los parámetros de tag y referencia en el script localizador, obteniendo poses métricas coherentes en el plano de trabajo.",
            "Presenta errores de escala por mala medición de tag_size_m o desorientación de los ejes del marco de referencia.",
            "No implementa la homografía, confunde marcos de coordenadas o carece de justificación de interfaces.",
        ],
    },
    {
        "code": "C3",
        "name": "Exactitud de localización, paralaje y telemetría",
        "weight": 25,
        "so": "Student Outcome SO6",
        "indicator": INDICATOR_61,
        "descriptors": [
            "Además de N4, modela matemáticamente el error de paralaje en función de la altura del vehículo y el ángulo visual de la cámara, proponiendo una compensación geométrica directa en el localizador.",
            "Diligencia completamente la Tabla 2 con 5 puntos de prueba, cuantifica el error en plano vs techo del carrito, mide las frecuencias del lazo (Tabla 3) y valida el LED fijo a < 5 cm.",
            "Realiza mediciones de exactitud en al menos 3 posiciones, verifica el cambio de estado del LED y registra las tasas de los tópicos.",
            "Mediciones incompletas en la Tabla 2, omisión de la evaluación de paralaje o tasas de tópicos inestables sin justificación.",
            "No realiza mediciones experimentales o afirma exactitud sin datos verificables.",
        ],
    },
    {
        "code": "C4",
        "name": "Robustez del lazo, oclusión y resiliencia",
        "weight": 15,
        "so": "Student Outcome SO6",
        "indicator": INDICATOR_64,
        "descriptors": [
            "Además de N4, implementa y valida en el firmware el mini-reto del Watchdog de pose con millis(), demostrando apagado seguro automático ante oclusiones > 500 ms y reconexión inmediata.",
            "Ejecuta las pruebas de oclusión de tags (Tabla 4), explica el síntoma de pose congelada, mide el tiempo de reconexión autónoma del agente y formula la solución por timeout.",
            "Prueba la oclusión de tags y la desconexión del agente, documentando las respuestas observadas en la Tabla 4.",
            "Confunde la oclusión con desconexión de red o requiere reset físico forzado del microcontrolador para recuperar el lazo.",
            "No ejecuta pruebas de robustez o no presenta evidencia de manejo de fallos.",
        ],
    },
    {
        "code": "C5",
        "name": "Trazabilidad en rosbag MCAP y trabajo en equipo",
        "weight": 10,
        "so": "Student Outcome SO3 / SO5",
        "indicator": f"{INDICATOR_31_33} / {INDICATOR_51}",
        "descriptors": [
            "Además de N4, presenta un script de extracción programática (rosbag2_py) que grafica la correlación temporal y el desfase exacto entre pose y distancia, con sustentación individual sobresaliente en el Anexo A.",
            "Graba el dataset MCAP con metadatos válidos, genera gráficas claras de evolución temporal (x, y, d), entrega el informe técnico estructurado y el Anexo A evidencia autoría individual.",
            "Entrega el dataset grabado con ros2 bag info, informe con tablas llenas y Anexo A con respuestas verificables.",
            "Dataset con pérdida de mensajes, tablas incompletas o Anexo A con respuestas genéricas sin evidencia de autoría.",
            "No entrega rosbag ni informe, o el Anexo A demuestra ausencia de participación individual.",
        ],
    },
]


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tc_pr.append(node)
    node.set(qn("w:fill"), fill)


def format_cell(cell, *, bold: bool = False, size: float = 8, color: str | None = None, align: WD_ALIGN_PARAGRAPH | None = None) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        if align is not None:
            paragraph.alignment = align
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        for run in paragraph.runs:
            run.font.name = FONT
            run.font.size = Pt(size)
            run.bold = bold
            if color:
                run.font.color.rgb = RGBColor.from_string(color)


def style_table(table, *, header: bool = True, font_size: float = 8) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            if header and i == 0:
                shade(cell, BLUE)
                format_cell(cell, bold=True, size=font_size, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
            else:
                if i % 2 == 0:
                    shade(cell, LIGHT_GRAY)
                format_cell(cell, size=font_size)


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = FONT
        run.font.color.rgb = RGBColor.from_string(BLUE)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    return p


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths=None, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    for j, value in enumerate(headers):
        table.cell(0, j).text = str(value)
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j, value in enumerate(row):
            cells[j].text = str(value)
    style_table(table, font_size=7.5)
    if widths:
        for row in table.rows:
            for j, width in enumerate(widths):
                if j < len(row.cells):
                    row.cells[j].width = Cm(width)
    if alignments:
        for row in table.rows[1:]:
            for j, align in enumerate(alignments):
                if j < len(row.cells):
                    row.cells[j].paragraphs[0].alignment = align
    return table


def add_note(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    shade(table.cell(0, 0), LIGHT_BLUE)
    table.cell(0, 0).text = text
    format_cell(table.cell(0, 0), size=8)


def add_response_box(doc: Document, prompt: str, lines: int = 4) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(prompt)
    run.font.name = FONT
    run.font.size = Pt(8.5)
    run.font.italic = True

    box = doc.add_table(rows=1, cols=1)
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    box.style = "Table Grid"
    cell = box.cell(0, 0)
    shade(cell, "FAFAFA")
    cell.width = Cm(17.5)
    cp = cell.paragraphs[0]
    cp.paragraph_format.space_before = Pt(2)
    cp.paragraph_format.space_after = Pt(2)
    r = cp.add_run("Respuesta Técnica del Equipo:\n" + ("\n" * (lines - 1)))
    r.font.name = FONT
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string("555555")


def set_doc_defaults(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    styles = doc.styles
    styles["Normal"].font.name = FONT
    styles["Normal"].font.size = Pt(9)


def add_page_number_footer(doc: Document) -> None:
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("ROB-ROS · Instrumento Único de Evidencia y Evaluación ABET — Taller AprilTag micro-ROS · 2026-2")
    run.font.name = FONT
    run.font.size = Pt(8)


def build_instrument():
    doc = Document()
    set_doc_defaults(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(
        "INSTRUMENTO DE EVIDENCIA Y CALIFICACIÓN ACADÉMICA\n"
        "TALLER: LOCALIZACIÓN VISUAL 2D CON APRILTAGS (KINOVA)\n"
        "Y LAZO DE POSE DISTRIBUIDO CON MICRO-ROS EN ESP32"
    )
    run.bold = True
    run.font.name = FONT
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    add_note(
        doc,
        "DOCUMENTO ÚNICO DE ENTREGA Y EVALUACIÓN ABET — Versión 1.0 (2026-2). "
        "Asociado al Taller de Localización Visual y micro-ROS. Este instrumento contiene el informe técnico "
        "integral del equipo: las 4 tablas de registro experimental diligenciadas (Modos A y B), el cuestionario "
        "de análisis técnico respondido, el Anexo A de comprobación individual para certificar el logro de cada estudiante "
        "y la rúbrica analítica docente para calificación sobre 500 y 5,0.",
    )

    # 1. Identificación
    add_heading(doc, "1. Identificación y Control")
    add_table(
        doc,
        ["Campo", "Registro Oficial"],
        [
            ["Programa Académico", "Ingeniería Mecatrónica"],
            ["Asignatura", "ROBOT OPERATING SYSTEM - ROS"],
            ["Periodo Académico", "2026-2"],
            ["Corte / Instrumento", "Primer Corte / Talleres y tareas"],
            ["Actividad Evaluada", "Taller — Localización Visual 2D con AprilTags (Kinova) y micro-ROS en ESP32"],
            ["Número de Grupo / Subgrupo", ""],
            ["Estudiante 1 (Nombre y Código)", ""],
            ["Estudiante 2 (Nombre y Código)", ""],
            ["Estudiante 3 (si aplica)", ""],
            ["Namespace Asignado al Carrito", "/burger_car_"],
            ["Modos Evaluados", "Modo A (Simulado) ☐        Modo B (Real Kinova) ☐"],
            ["Archivo de Entrega", "C1_T_APRILTAG_G<grupo>_<codigo1>_<codigo2>_v1.docx"],
            ["Fecha de Realización en Laboratorio", ""],
            ["Fecha de Entrega del Documento", ""],
            ["Docente Evaluador", "Ing. Henry Roncancio"],
            ["Versión del Instrumento", "Versión 1.0 (2026-2)"],
            ["Unidad de Análisis / Captura", "Equipo colaborativo con comprobación individual (Anexo A)"],
        ],
        [5, 12],
    )

    # 2. Parámetros de assessment
    add_heading(doc, "2. Parámetros de Assessment")
    add_table(
        doc,
        ["Parámetro", "Regla Institucional y Metodológica Adoptada"],
        [
            ["Población o cohorte", "Censo completo de estudiantes matriculados que presentan el taller en el periodo 2026-2."],
            ["Momento de medición", "Primer corte, tras la ejecución de las fases de visión, homografía y lazo embebido."],
            ["Evaluador", "Docente titular de la asignatura ROBOT OPERATING SYSTEM - ROS."],
            ["Umbral individual de logro", "Nivel N3 o superior (puntaje mínimo de 300 sobre 500) en cada indicador evaluado."],
            ["Meta de cohorte", "Al menos el 70% de los estudiantes evaluables debe alcanzar el nivel N3 o superior en cada indicador."],
            ["Regla de muestreo", "Censo 100%: no se utiliza muestreo; se evalúa la totalidad de los estudiantes y equipos exigibles."],
            ["Evidencia faltante", "Entrega exigible sin evidencia obligatoria verificable: N1 con valor 0. Retiro oficial: NA / no evaluado."],
            ["Regla de trabajo colaborativo", "Informe único por equipo con tablas llenas más Anexo A individual obligatorio."],
        ],
        [5, 12],
    )

    # 3. Niveles Zubatronic
    add_heading(doc, "3. Niveles de Desempeño para Zubatronic/SGDE")
    add_table(doc, ["Nivel", "Intervalo", "Interpretación y Criterio de Logro"], [list(x) for x in LEVELS], [2, 3, 12])

    # 4. Alineación Criterios
    add_heading(doc, "4. Alineación de Criterios, RAE y Student Outcomes")
    add_table(
        doc,
        ["Criterio", "Peso", "Student Outcome Principal", "Indicador Literal del Programa", "Evidencia Directa Obligatoria"],
        [
            [
                f"{c['code']}. {c['name']}",
                f"{c['weight']}%",
                c["so"],
                c["indicator"],
                "rqt_graph, micro_ros_agent, Subscription count: 1, Tabla 1" if c['code'] == 'C1' else
                ("Homografía tag_mesa, tag_size_m pie de rey, Pose2D, Tabla 1" if c['code'] == 'C2' else
                ("Tabla 2 exactitud/paralaje (5 pts), ros2 topic hz (Tabla 3), LED" if c['code'] == 'C3' else
                ("Oclusión de tags, congelamiento de LED, timeout millis, Tabla 4" if c['code'] == 'C4' else
                "Dataset MCAP, gráficas x,y,d vs tiempo, Anexo A individual")))
            ]
            for c in CRITERIA
        ],
        [4.2, 1.3, 2.5, 5.0, 4.5],
    )

    # 5. Registro de Evidencias
    add_heading(doc, "5. Registro de Evidencias de la Práctica (E1–E8)")
    add_table(
        doc,
        ["Código", "Evidencia Requerida", "Localizador en Documento / Repositorio / Archivo"],
        [
            ["E1", "Captura de rqt_graph en Modo B mostrando /camera/kinova_vision_color, localizador y nodo ESP32.", ""],
            ["E2", "Configuración de tags: IDs, medición con pie de rey de tag_size_m, ejes de tag_mesa y cálculo a mano Ej 3.1.", ""],
            ["E3", "Captura de terminal con ros2 topic info /burger_car_NN/pose2d mostrando Subscription count: 1 (ESP32).", ""],
            ["E4", "Captura de ros2 topic echo de distance_to_goal convergiendo a < 0.05 m y fotografía del LED encendido fijo.", ""],
            ["E5", "Medición de tasas con ros2 topic hz para imagen comprimida, pose2d y distance_to_goal.", ""],
            ["E6", "Captura del comportamiento del lazo ante oclusión del tag del carrito y de tag_mesa (congelamiento).", ""],
            ["E7", "Cronometraje y captura de reconexión autónoma del ESP32 tras detener y relanzar el agente UDP.", ""],
            ["E8", "Salida de ros2 bag info dataset_lazo_burger_car_NN y gráfica de x, y, distancia vs tiempo de recepción.", ""],
        ],
        [2, 11, 4.5],
    )

    # 6. Tablas Experimentales Llenas
    add_heading(doc, "6. Tablas de Registro Experimental (Diligenciadas por el Equipo)")
    add_note(doc, "Instrucciones: Registre los datos cuantitativos obtenidos durante la sesión de laboratorio en Modo A y Modo B.")

    p = doc.add_paragraph("Tabla 1: Configuración de Tags, Dispositivos y Parámetros del Localizador")
    p.runs[0].bold = True
    add_table(
        doc,
        ["Parámetro", "Valor Configurado / Medido", "Método de Verificación"],
        [
            ["Familia de AprilTags", "Tag36h11", "Inspección visual"],
            ["ID del Tag de Referencia (tag_mesa)", "", "Parámetro -p reference_tag_id:="],
            ["ID del Tag del Carrito (tag_id)", "", "Parámetro -p tag_id:="],
            ["Tamaño real del cuadro negro (tag_size_m)", "m (ej. 0.100 m)", "Medición con pie de rey"],
            ["Resolución del flujo visual", "píxeles", "ros2 topic echo /camera/color/camera_info --once"],
            ["Namespace del robot móvil", "/burger_car_", "Parámetro -p robot_namespace:="],
            ["Dirección IP y puerto del Agente micro-ROS", "UDP puerto 8888", "Configuración en PC"],
            ["Dirección IP asignada al ESP32", "", "Monitor Serie ('IP del ESP32: ...')"],
        ],
        [6.0, 5.5, 6.0],
    )

    p = doc.add_paragraph("Tabla 2: Exactitud de Localización 2D y Medición de Paralaje (Modo B)")
    p.runs[0].bold = True
    add_table(
        doc,
        ["Punto", "Coord Real Mesa (x, y) [m]", "Medición Plano (x, y) [m]", "Error Plano [m]", "Medición Carrito (x, y) [m]", "Error Carrito [m]", "Error Paralaje [m]"],
        [
            ["P1", "(0.00, 0.00)", "", "", "", "", ""],
            ["P2", "(0.20, 0.15)", "", "", "", "", ""],
            ["P3", "(0.40, 0.00)", "", "", "", "", ""],
            ["P4", "(0.00, 0.30)", "", "", "", "", ""],
            ["P5", "(0.35, 0.25)", "", "", "", "", ""],
        ],
        [1.5, 3.5, 3.2, 2.3, 3.2, 2.3, 2.2],
    )

    p = doc.add_paragraph("Tabla 3: Medición de Tasas de Frecuencia y Sincronismo del Lazo Distribuido")
    p.runs[0].bold = True
    add_table(
        doc,
        ["Tópico del Lazo", "Tasa Nominal / Esperada", "Tasa Media Medida (Hz)", "Mínimo (s)", "Máximo (s)", "Desviación Estándar (std dev)"],
        [
            ["/camera/color/image_raw/compressed", "15–30 Hz", "Hz", "", "", ""],
            ["/burger_car_NN/pose2d", "Igual a la imagen", "Hz", "", "", ""],
            ["/burger_car_NN/distance_to_goal", "Igual a la pose", "Hz", "", "", ""],
        ],
        [5.5, 2.5, 2.5, 2.2, 2.2, 2.6],
    )

    p = doc.add_paragraph("Tabla 4: Pruebas de Robustez, Oclusión y Resiliencia del Enlace")
    p.runs[0].bold = True
    add_table(
        doc,
        ["Escenario Experimental", "Acción Ejecutada", "Respuesta del Localizador", "Respuesta ESP32 y LED", "Tiempo Recuperación"],
        [
            ["Oclusión Tag Carrito", "Tapar tag del carrito 10 s", "Avisa no visible, cesa /pose2d", "LED se congela", ""],
            ["Oclusión Tag Mesa", "Tapar tag_mesa 10 s", "Avisa ref no visible, cesa /pose2d", "LED se congela", ""],
            ["Caída Agente micro-ROS", "Ctrl+C al agente en PC", "Continúa publicando /pose2d", "LED se apaga / desconexión", ""],
            ["Reinicio de Agente", "Relanzar agente UDP", "Normal", "Reconexión autónoma", "s (sin reset físico)"],
        ],
        [3.0, 3.5, 4.0, 4.0, 3.0],
    )

    # 7. Cuestionario Técnico
    add_heading(doc, "7. Cuestionario de Análisis Técnico (Respondido por el Equipo)")

    add_response_box(
        doc,
        "Pregunta 1: Ausencia del Agente micro-ROS en rqt_graph\n"
        "¿Por qué el nodo micro_ros_agent no aparece como un nodo visible en rqt_graph ni en ros2 node list, "
        "mientras que el nodo embebido /burger_car_NN/visual_navigator sí aparece en el grafo DDS a pesar de "
        "estar ejecutándose físicamente dentro del ESP32? Explique el rol de puente XRCE-DDS ↔ DDS.",
        lines=4,
    )

    add_response_box(
        doc,
        "Pregunta 2: Justificación de geometry_msgs/msg/Pose2D y Ausencia de Header\n"
        "Justifique por qué en este lazo embebido se seleccionó el mensaje compacto geometry_msgs/msg/Pose2D "
        "en lugar de geometry_msgs/msg/PoseStamped. ¿Qué ventajas aporta a la memoria y ancho de banda del "
        "ESP32 y qué desventaja crítica introduce la ausencia del campo std_msgs/Header?",
        lines=4,
    )

    add_response_box(
        doc,
        "Pregunta 3: Origen Físico y Mitigación del Error de Paralaje\n"
        "A partir de los datos registrados en la Tabla 2, explique físicamente por qué la estimación de la pose "
        "del carrito difiere cuando el tag está colocado sobre la mesa frente a cuando está en el techo del vehículo "
        "(z > 0). Si la homografía asume z=0, ¿cómo influye el ángulo de inclinación de la cámara en este error?",
        lines=4,
    )

    add_response_box(
        doc,
        "Pregunta 4: Diagnóstico del LED Congelado y Necesidad de Timeout de Pose\n"
        "Durante la prueba de oclusión (Ejercicio 6.2), el LED del ESP32 se congela en su último estado cuando se "
        "tapa el tag. Explique por qué un LED encendido fijo NO garantiza que el carrito esté realmente en la meta. "
        "Proponga cómo implementar un mecanismo de seguridad (Watchdog) en el firmware con millis().",
        lines=4,
    )

    add_response_box(
        doc,
        "Pregunta 5: Comparativa Metodológica: Modo A (Simulado) vs. Modo B (Real)\n"
        "Compare las fortalezas y limitaciones de validar el lazo en Modo A frente a Modo B. ¿Qué componentes de la "
        "cadena de ingeniería valida con certeza el Modo A y cuáles enmascara por completo? ¿Por qué es una buena práctica?",
        lines=4,
    )

    add_response_box(
        doc,
        "Pregunta 6: Uso de la Marca Temporal de Recepción en rosbag2\n"
        "En el Ejercicio 6.4 se grabó el dataset MCAP con ros2 bag record. Dado que Pose2D no incluye header.stamp, "
        "¿cómo permite rosbag2 reconstruir la relación temporal y latencia entre la pose y la distancia calculada?",
        lines=4,
    )

    # 8. Anexo A
    add_heading(doc, "8. Anexo A — Comprobación Individual de Desempeño")
    add_note(doc, "Propósito ABET: Verifica la autoría, comprensión técnica y contribución directa individual de cada estudiante.")

    for est_num in [1, 2]:
        p = doc.add_paragraph(f"Estudiante {est_num}: __________________________________________________ Código: ____________________")
        p.runs[0].bold = True
        add_table(
            doc,
            ["Pregunta Individual de Verificación", "Respuesta / Evidencia Directa del Estudiante"],
            [
                ["1. Rol y tareas técnicas desarrolladas: Describa sus responsabilidades específicas en la práctica (ej. OpenCV, medición de tags, agente, firmware ESP32 o rosbag).", ""],
                ["2. Comprensión de la homografía: Explique qué hace la homografía con las esquinas de tag_mesa y cómo transforma píxeles a metros en el plano de trabajo.", ""],
                ["3. Cadena de transporte XRCE-DDS: Explique cómo viaja el mensaje desde que OpenCV calcula la pose hasta que la callback del ESP32 recibe los datos y actualiza el LED.", ""],
                ["4. Análisis de una anomalía resuelta: Describa un error o fallo enfrentado durante la práctica (ej. iluminación, tags no detectados, Wi-Fi) y cómo lo resolvió.", ""],
                ["5. Autoría y reproducibilidad: Indique qué secciones del código, tablas, capturas o gráficas del informe fueron elaboradas directamente por usted.", ""],
            ],
            [7.5, 10.0],
        )
        p = doc.add_paragraph(f"Estado de comprobación individual Estudiante {est_num}: Verificada ☐     Insuficiente ☐     NA autorizado ☐     Firma: ________________________")
        p.paragraph_format.space_after = Pt(6)

    # 9. Rúbricas Docentes
    add_heading(doc, "9. Selección del Nivel Alcanzado por Criterio (Rúbricas Docentes)")
    for criterion in CRITERIA:
        add_heading(
            doc,
            f"{criterion['code']}. {criterion['name']} — {criterion['weight']}% — {criterion['so']}",
            level=2,
        )
        rows = []
        for level, descriptor in zip(LEVELS, criterion["descriptors"]):
            rows.append(["☐", f"{level[0]} — {level[1]}", descriptor])
        add_table(doc, ["Marque", "Nivel", "Evidencia observable"], rows, [1.5, 3, 12.5])
        doc.add_paragraph(f"Nivel {criterion['code']} marcado: ________     Valor Zubatronic (0–500): ________")

    # 10. Consolidado Final
    add_heading(doc, "10. Consolidado Final de Calificación Académica")
    add_table(
        doc,
        ["Criterio", "Peso Oficial (%)", "Nivel Marcado", "Valor Obtenido (0–500)", "Aporte Ponderado"],
        [
            [f"{c['code']}. {c['name']}", f"{c['weight']}%", "", "", ""] for c in CRITERIA
        ] + [["TOTAL CONSOLIDADO", "100%", "", "", "________ / 500"]],
        [6.5, 2.5, 2.2, 3.3, 3.0],
    )
    add_note(
        doc,
        "Nota Académica sobre 5,0 = Nota Consolidada sobre 500 ÷ 100. "
        "Aporte a Talleres del Primer Corte: T₁ = Nota Académica sobre 5,0.",
    )

    add_table(
        doc,
        ["Resultado Oficial de la Actividad", "Registro Oficial"],
        [
            ["Nota de Taller AprilTag sobre 500 puntos", "__________ / 500"],
            ["Nota Académica Oficial sobre 5,0", "__________ / 5,0"],
            ["Número de Criterios en Nivel N3 o superior (Umbral Individual)", "_____ / 5"],
            ["¿Cumple el Umbral Individual de Logro ABET (todos en N3+)?", "SÍ ☐         NO ☐"],
        ],
        [8.5, 9.0],
    )

    # 11. Consolidación ABET
    add_heading(doc, "11. Consolidación ABET y Cierre de Mejora Continua")
    add_table(
        doc,
        ["Student Outcome / Indicador Evaluado", "N Evaluables", "N en N3+", "% Logro", "Meta", "Hallazgo Docente"],
        [
            ["SO1 / Indicador 1.1 (Resolución de Problemas y Lazo)", "", "", "", "70%", ""],
            ["SO2 / Indicador 2.2 (Diseño e Interfaces Compactas)", "", "", "", "70%", ""],
            ["SO6 / Indicador 6.1 y 6.4 (Experimentación y Diagnóstico)", "", "", "", "70%", ""],
            ["SO3 / Indicador 3.1 y 3.3 (Documentación y Rosbag)", "", "", "", "70%", ""],
            ["SO5 / Indicador 5.1 (Trabajo en Equipo y Roles)", "", "", "", "70%", ""],
        ],
        [5.0, 2.0, 2.0, 2.0, 1.8, 4.7],
    )

    add_table(
        doc,
        ["Campo de Cierre de Ciclo de Mejora Continua", "Registro de Gestión Pedagógica"],
        [
            ["Decisión derivada del hallazgo", ""],
            ["Acción de mejora continua pedagógica y técnica", ""],
            ["Responsable de la acción", "Ing. Henry Roncancio"],
            ["Fecha prevista de seguimiento", ""],
            ["Evidencia de seguimiento", ""],
            ["Resultado observado en el segundo ciclo", ""],
        ],
        [6.0, 11.5],
    )

    # 12. Cierre y Firmas
    add_heading(doc, "12. Cierre y Firmas de Conformidad")
    doc.add_paragraph("Docente Evaluador: Ing. Henry Roncancio           Fecha de Evaluación: ____________________")
    doc.add_paragraph("Firma del Docente Evaluador: __________________________________________________")
    doc.add_paragraph("Observaciones Finales y Retroalimentación Pedagógica:\n________________________________________________________________________________\n________________________________________________________________________________")

    add_page_number_footer(doc)
    doc.core_properties.title = "Instrumento ABET Taller AprilTag y micro-ROS"
    doc.core_properties.subject = "ROBOT OPERATING SYSTEM - ROS · 2026-2"
    doc.core_properties.author = "Ing. Henry Roncancio"
    doc.core_properties.comments = "Versión 1.0. Documento único de entrega por equipos y evaluación ABET con tablas de exactitud, preguntas técnicas y Anexo A."
    TARGET.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(TARGET))
    print(f"Instrumento ABET Word generado exitosamente en: {TARGET}")


if __name__ == "__main__":
    build_instrument()
