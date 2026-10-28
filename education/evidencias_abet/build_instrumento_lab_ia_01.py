#!/usr/bin/env python3
"""
Constructor para INSTRUMENTO_ABET_LAB_IA_01_EMBODIED_AI_GEMINI_ER2.docx

Convierte la fuente INSTRUMENTO_ABET_LAB_IA_01_EMBODIED_AI_GEMINI_ER2.md en el documento
Word único de entrega y evaluación ABET del Laboratorio IA 01 (Embodied AI con Gemini
Robotics-ER 2), con el mismo estilo que los instrumentos de los laboratorios de ROS.

Reutiliza el analizador de Markdown del constructor de la guía IA 01 y las utilidades de
formato del instrumento del Laboratorio 03, para que el Markdown sea la única fuente.
Ejecución:  python3 education/evidencias_abet/build_instrumento_lab_ia_01.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

AQUI = Path(__file__).resolve().parent
sys.path.insert(0, str(AQUI))
sys.path.insert(0, str(AQUI.parent / "guias_laboratorio"))
import build_guia_lab_ia_01_embodied as md  # noqa: E402
import build_instrumento_lab_03 as estilo  # noqa: E402
from build_guia_lab_02_camara_kinova import clean  # noqa: E402

FUENTE = AQUI / "INSTRUMENTO_ABET_LAB_IA_01_EMBODIED_AI_GEMINI_ER2.md"
DESTINO = AQUI / "INSTRUMENTO_ABET_LAB_IA_01_EMBODIED_AI_GEMINI_ER2.docx"
PIE = "IA · Instrumento Único de Evidencia y Evaluación ABET — Lab IA 01 · 2026-2"
# Metadatos del .docx; otros instrumentos generados desde Markdown los cambian
# (build_instrumento_lab_04.py).
TITULO_DOC = "Instrumento ABET – Laboratorio IA 01 – Embodied AI con Gemini Robotics-ER 2"
ASIGNATURA = "INTELIGENCIA ARTIFICIAL"
CONSTRUCTOR = Path(__file__).name


def plano(texto: str) -> str:
    return clean(md.latex_a_texto(texto.replace("&nbsp;", " ")))


def parrafo(doc, texto: str, *, size: float = 9, italic: bool = False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(plano(texto))
    r.font.name, r.font.size, r.italic = estilo.FONT, Pt(size), italic
    return p


def construir():
    doc = Document()
    estilo.set_doc_defaults(doc)

    for tipo, datos in md.bloques(FUENTE.read_text(encoding="utf-8")):
        if tipo == "h1":
            titulo = doc.add_paragraph()
            titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = titulo.add_run("INSTRUMENTO DE EVIDENCIA Y CALIFICACIÓN ACADÉMICA\n"
                               + plano(datos.split("—", 1)[-1]).upper())
            r.bold, r.font.name, r.font.size = True, estilo.FONT, Pt(13)
            r.font.color.rgb = RGBColor.from_string(estilo.BLUE)
        elif tipo == "h2":
            estilo.add_heading(doc, plano(datos), level=1)
        elif tipo in ("h3", "h4"):
            estilo.add_heading(doc, plano(datos), level=2)
        elif tipo == "tabla":
            filas = [[plano(c) for c in fila] for fila in datos]
            tabla = estilo.add_table(doc, filas[0], filas[1:])
            md.fijar_anchos(tabla, md.anchos(datos, total_cm=17.5))
            if filas[0][0] == "Marque":  # rúbricas: casillas centradas
                for fila in tabla.rows[1:]:
                    fila.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif tipo == "aviso":
            _, texto = datos
            if "Respuesta Técnica del Equipo" in texto:
                estilo.add_response_box(doc, "", lines=6)
            else:
                estilo.add_note(doc, plano(texto))
        elif tipo == "codigo":
            p = doc.add_paragraph()
            r = p.add_run("\n".join(datos))
            r.font.name, r.font.size = "Courier New", Pt(8.5)
        elif tipo == "lista":
            for _, num, texto in datos:
                parrafo(doc, (f"{num}. " if num else "• ") + texto)
        elif tipo == "parrafo":
            if datos.startswith("*") and datos.endswith("*") and not datos.startswith("**"):
                parrafo(doc, datos, size=8.5, italic=True)  # enunciado de pregunta
            elif datos.startswith("**Observaciones"):
                estilo.add_response_box(doc, plano(datos.split("____")[0]), lines=3)
            else:
                parrafo(doc, datos)

    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(PIE)
    r.font.name, r.font.size = estilo.FONT, Pt(8)

    cp = doc.core_properties
    cp.title = TITULO_DOC
    cp.subject, cp.author = ASIGNATURA, "Universidad Militar Nueva Granada"
    cp.comments = f"Generado desde {FUENTE.name} con {CONSTRUCTOR}."
    doc.save(DESTINO)
    print(f"Instrumento ABET Word generado en: {DESTINO}")


if __name__ == "__main__":
    construir()
