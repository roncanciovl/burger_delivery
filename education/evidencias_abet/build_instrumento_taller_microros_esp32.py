#!/usr/bin/env python3
"""
Constructor para INSTRUMENTO_ABET_TALLER_MICROROS_ESP32_ROBOTICA_MOVIL.docx
Genera el instrumento Word único de captura de evidencia, tablas de datos experimentales,
cuestionario de análisis técnico, evaluación y consolidación ABET para el Taller de
Integración de micro-ROS en ESP32 para Plataformas Móviles y Drones (Serial y WiFi UDP).
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
TARGET = ROOT / "education" / "evidencias_abet" / "INSTRUMENTO_ABET_TALLER_MICROROS_ESP32_ROBOTICA_MOVIL.docx"

FONT = "Arial"
BLUE = "1F4E79"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F2F2"

INDICATOR_11 = (
    "Indicador de desempeño 1.1: Identifica y selecciona los requerimientos para la arquitectura "
    "de software distribuido del robot mediante nodos de ROS 2 y micro-ROS."
)
INDICATOR_22 = (
    "Indicador de desempeño 2.2: Incorpora restricciones de red, latencia, ancho de banda y seguridad "
    "en la integración de hardware heterogéneo (micro-ROS en microcontroladores)."
)
INDICATOR_61 = (
    "Indicador de desempeño 6.1: Diseña y ejecuta pruebas de conectividad, jitter, pérdida de paquetes "
    "y latencia en enlaces inalámbricos XRCE-DDS."
)
INDICATOR_64 = (
    "Indicador de desempeño 6.4: Interpreta fallas y diagnósticos experimentales aplicando protocolos "
    "de diagnóstico por capas para aislar errores en hardware y software."
)
INDICATOR_31_33 = (
    "Indicador de desempeño 3.1 - 3.3: Elabora documentación técnica reproducible del sistema ROS 2 "
    "y comunica resultados experimentales de plataformas móviles y telemetría."
)
INDICATOR_51 = (
    "Indicador de desempeño 5.1: Define roles técnicos, coordina la ejecución en equipo y fomenta "
    "un entorno de trabajo colaborativo e inclusivo."
)

LEVELS = [
    ("N5", "475–500", "Excelente: solución y evidencia completas, precisas, reproducibles y explicadas con profundidad analítica y técnica sobresaliente."),
    ("N4", "400–474", "Bueno: desempeño técnico correcto con omisiones menores que no comprometen la operación, reproducibilidad ni diagnóstico."),
    ("N3", "300–399", "Aceptable: demuestra el desempeño esencial con evidencia verificable. Es el umbral individual de logro."),
    ("N2", "150–299", "Cumplimiento parcial: evidencia incompleta, métricas faltantes o errores conceptuales en XRCE-DDS o redes."),
    ("N1", "0–149", "No cumple: evidencia mínima, fragmentaria o no funcional. Sin evidencia obligatoria se registra 0."),
]

CRITERIA = [
    {
        "code": "C1",
        "name": "Arquitectura Micro XRCE-DDS y agentes micro-ROS",
        "weight": 25,
        "so": "Student Outcome SO1",
        "indicator": INDICATOR_11,
        "descriptors": [
            "Además de N4, analiza la serialización CDR en tramas XRCE-DDS, optimiza el tamaño de los buffers del middleware y compara el overhead de transporte serial vs UDP demostrando uso eficiente de la memoria estática.",
            "Configura y ejecuta exitosamente el agente micro-ROS en PC (Docker o nativo), verifica la creación de clientes y sesiones con -v6, y valida la comunicación tanto por puerto Serial como por Wi-Fi UDP.",
            "Pone en marcha el agente micro-ROS, establece comunicación con el ESP32 y documenta la salida en las Tablas 1 y 2.",
            "Presenta fallos de conexión por configuración de puertos, discrepancia de bibliotecas o no logra verificar el flujo serial.",
            "No logra ejecutar el agente micro-ROS o carece de evidencias funcionales obligatorias.",
        ],
    },
    {
        "code": "C2",
        "name": "Namespaces, interfaces y teleoperación en tiempo real",
        "weight": 25,
        "so": "Student Outcome SO2",
        "indicator": INDICATOR_22,
        "descriptors": [
            "Además de N4, implementa el Mini-Reto 3 accionando motores reales mediante PWM con saturación segura y parada automática de hardware programada en destroy_entities().",
            "Configura el namespace /burger_car_NN, valida que los tópicos sean relativos, verifica Subscription count: 1 en cmd_vel, y demuestra teleoperación interactiva con teleop_twist_keyboard y respuesta en LED.",
            "Asigna namespace al nodo, recibe comandos de velocidad encendiendo el LED y visualiza la telemetría sintética de batería.",
            "Opera con tópicos en el espacio global sin namespace causando colisiones o requiere modificar código para cambiar metas.",
            "No logra recibir comandos de velocidad en el microcontrolador o carece de justificación de interfaces.",
        ],
    },
    {
        "code": "C3",
        "name": "Diagnóstico por tramos de red y métricas cuantitativas",
        "weight": 25,
        "so": "Student Outcome SO6",
        "indicator": INDICATOR_61,
        "descriptors": [
            "Además de N4, analiza la distribución estadística del jitter (mdev), correlaciona las fluctuaciones de latencia con la saturación del espectro Wi-Fi y propone umbrales de QoS para navegación autónoma en interiores.",
            "Ejecuta el script diagnostico_microros.sh, mide el ping de 100 paquetes cerca y lejos del router (Tabla 3), analiza la estabilidad de telemetría a 1, 5 y 10 Hz (Tabla 4) y diferencia claramente los tramos de red.",
            "Realiza mediciones de ping hacia el ESP32, reporta latencias y registra la estabilidad de telemetría a dos frecuencias distintas.",
            "Mediciones incompletas en las Tablas 3 y 4 o confusión entre las mediciones del monitor web (router) y las del ESP32.",
            "No realiza mediciones cuantitativas de red o reporta valores ficticios sin captura de terminal verificable.",
        ],
    },
    {
        "code": "C4",
        "name": "Resiliencia, máquina de estados y reconexión",
        "weight": 15,
        "so": "Student Outcome SO6",
        "indicator": INDICATOR_64,
        "descriptors": [
            "Además de N4, instrumenta el firmware con telemetría de fallos, cuantifica el consumo de heap en reconexiones sucesivas demostrando cero fugas de memoria y formula un diagrama de estados formal.",
            "Demuestra la simetría entre create_entities() y destroy_entities(), cronometra la reconexión autónoma tras detener y relanzar el agente (Tabla 5) y verifica que no requiere reset físico.",
            "Realiza la prueba de apagado del agente, documenta la pérdida de sesión y logra la reconexión tras reiniciar el agente.",
            "El ESP32 se bloquea al perder la conexión con el agente, requiere reset por pulsador físico o presenta fugas de memoria.",
            "No implementa máquina de estados de reconexión o carece de evidencias de resiliencia.",
        ],
    },
    {
        "code": "C5",
        "name": "Trazabilidad técnica y trabajo en equipo",
        "weight": 10,
        "so": "Student Outcome SO3 / SO5",
        "indicator": f"{INDICATOR_31_33} / {INDICATOR_51}",
        "descriptors": [
            "Además de N4, el informe técnico incluye diagramas de secuencia UML de la negociación XRCE-DDS y el Anexo A demuestra una división de roles técnicos y dominio individual sobresaliente de ambos integrantes.",
            "Presenta el informe técnico completo con comandos reproducibles, capturas de rqt_graph comentadas, tablas llenas con rigor y el Anexo A evidencia autoría individual.",
            "Entrega el informe con tablas diligenciadas, capturas de terminal legibles y Anexo A con respuestas verificables.",
            "Informe incompleto, capturas sin contexto, tablas con datos faltantes o Anexo A con respuestas genéricas.",
            "No entrega el informe técnico o el Anexo A demuestra ausencia de participación individual.",
        ],
    },
]


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tc_pr.append(node)
    node.set(qn("w:fill"), fill)


def format_cell(cell, *, bold: bool = False, size: float = 8, color: str | None = None, align: WD_ALIGN_PARAGRAPH | None = None) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        if align is not None:
            paragraph.alignment = align
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        for run in paragraph.runs:
            run.font.name = FONT
            run.font.size = Pt(size)
            run.bold = bold
            if color:
                run.font.color.rgb = RGBColor.from_string(color)


def style_table(table, *, header: bool = True, font_size: float = 8) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            if header and i == 0:
                shade(cell, BLUE)
                format_cell(cell, bold=True, size=font_size, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
            else:
                if i % 2 == 0:
                    shade(cell, LIGHT_GRAY)
                format_cell(cell, size=font_size)


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = FONT
        run.font.color.rgb = RGBColor.from_string(BLUE)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    return p


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths=None, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    for j, value in enumerate(headers):
        table.cell(0, j).text = str(value)
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j, value in enumerate(row):
            cells[j].text = str(value)
    style_table(table, font_size=7.5)
    if widths:
        for row in table.rows:
            for j, width in enumerate(widths):
                if j < len(row.cells):
                    row.cells[j].width = Cm(width)
    if alignments:
        for row in table.rows[1:]:
            for j, align in enumerate(alignments):
                if j < len(row.cells):
                    row.cells[j].paragraphs[0].alignment = align
    return table


def add_note(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    shade(table.cell(0, 0), LIGHT_BLUE)
    table.cell(0, 0).text = text
    format_cell(table.cell(0, 0), size=8)


def add_response_box(doc: Document, prompt: str, lines: int = 4) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(prompt)
    run.font.name = FONT
    run.font.size = Pt(8.5)
    run.font.italic = True

    box = doc.add_table(rows=1, cols=1)
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    box.style = "Table Grid"
    cell = box.cell(0, 0)
    shade(cell, "FAFAFA")
    cell.width = Cm(17.5)
    cp = cell.paragraphs[0]
    cp.paragraph_format.space_before = Pt(2)
    cp.paragraph_format.space_after = Pt(2)
    r = cp.add_run("Respuesta Técnica del Equipo:\n" + ("\n" * (lines - 1)))
    r.font.name = FONT
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string("555555")


def set_doc_defaults(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    styles = doc.styles
    styles["Normal"].font.name = FONT
    styles["Normal"].font.size = Pt(9)


def add_page_number_footer(doc: Document) -> None:
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("ROB-ROS · Instrumento Único de Evidencia y Evaluación ABET — Taller micro-ROS ESP32 · 2026-2")
    run.font.name = FONT
    run.font.size = Pt(8)


def build_instrument():
    doc = Document()
    set_doc_defaults(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(
        "INSTRUMENTO DE EVIDENCIA Y CALIFICACIÓN ACADÉMICA\n"
        "TALLER: INTEGRACIÓN DE MICRO-ROS EN ESP32\n"
        "PARA PLATAFORMAS MÓVILES Y DRONES (SERIAL Y WIFI UDP)"
    )
    run.bold = True
    run.font.name = FONT
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    add_note(
        doc,
        "DOCUMENTO ÚNICO DE ENTREGA Y EVALUACIÓN ABET — Versión 1.0 (2026-2). "
        "Asociado al Taller de Integración de micro-ROS en ESP32. Este instrumento contiene el informe técnico "
        "integral del equipo: las 5 tablas de registro experimental diligenciadas (Serial y Wi-Fi UDP), el cuestionario "
        "de análisis técnico respondido, el Anexo A de comprobación individual para certificar el logro de cada estudiante "
        "y la rúbrica analítica docente para calificación sobre 500 y 5,0.",
    )

    # 1. Identificación
    add_heading(doc, "1. Identificación y Control")
    add_table(
        doc,
        ["Campo", "Registro Oficial"],
        [
            ["Programa Académico", "Ingeniería Mecatrónica"],
            ["Asignatura", "ROBOT OPERATING SYSTEM - ROS"],
            ["Periodo Académico", "2026-2"],
            ["Corte / Instrumento", "Segundo Corte / Talleres y tareas"],
            ["Actividad Evaluada", "Taller — Integración de micro-ROS en ESP32 (Transportes Serial y WiFi UDP, Namespace y Red)"],
            ["Número de Grupo / Subgrupo", ""],
            ["Estudiante 1 (Nombre y Código)", ""],
            ["Estudiante 2 (Nombre y Código)", ""],
            ["Estudiante 3 (si aplica)", ""],
            ["Namespace Asignado al Carrito", "/burger_car_"],
            ["Hardware Embebido Utilizado", "ESP32 NodeMCU ☐        ESP32-S3 ☐        ESP32-WROOM ☐"],
            ["Archivo de Entrega", "C2_T_MICROROS_ESP32_G<grupo>_<codigo1>_<codigo2>_v1.docx"],
            ["Fecha de Realización en Laboratorio", ""],
            ["Fecha de Entrega del Documento", ""],
            ["Docente Evaluador", "Ing. Henry Roncancio"],
            ["Versión del Instrumento", "Versión 1.0 (2026-2)"],
            ["Unidad de Análisis / Captura", "Equipo colaborativo con comprobación individual (Anexo A)"],
        ],
        [5, 12],
    )

    # 2. Parámetros de assessment
    add_heading(doc, "2. Parámetros de Assessment")
    add_table(
        doc,
        ["Parámetro", "Regla Institucional y Metodológica Adoptada"],
        [
            ["Población o cohorte", "Censo completo de estudiantes matriculados que presentan el taller en el periodo 2026-2."],
            ["Momento de medición", "Segundo corte, tras la ejecución de las fases de transporte serial, WiFi UDP, teleoperación y red."],
            ["Evaluador", "Docente titular de la asignatura ROBOT OPERATING SYSTEM - ROS."],
            ["Umbral individual de logro", "Nivel N3 o superior (puntaje mínimo de 300 sobre 500) en cada indicador evaluado."],
            ["Meta de cohorte", "Al menos el 70% de los estudiantes evaluables debe alcanzar el nivel N3 o superior en cada indicador."],
            ["Regla de muestreo", "Censo 100%: no se utiliza muestreo; se evalúa la totalidad de los estudiantes y equipos exigibles."],
            ["Evidencia faltante", "Entrega exigible sin evidencia obligatoria verificable: N1 con valor 0. Retiro oficial: NA / no evaluado."],
            ["Regla de trabajo colaborativo", "Informe único por equipo con tablas llenas más Anexo A individual obligatorio."],
        ],
        [5, 12],
    )

    # 3. Niveles Zubatronic
    add_heading(doc, "3. Niveles de Desempeño para Zubatronic/SGDE")
    add_table(doc, ["Nivel", "Intervalo", "Interpretación y Criterio de Logro"], [list(x) for x in LEVELS], [2, 3, 12])

    # 4. Alineación Criterios
    add_heading(doc, "4. Alineación de Criterios, RAE y Student Outcomes")
    add_table(
        doc,
        ["Criterio", "Peso", "Student Outcome Principal", "Indicador Literal del Programa", "Evidencia Directa Obligatoria"],
        [
            [
                f"{c['code']}. {c['name']}",
                f"{c['weight']}%",
                c["so"],
                c["indicator"],
                "micro_ros_agent (-v6), cliente/sesión, serial y UDP, Tablas 1 y 2" if c['code'] == 'C1' else
                ("Namespace /burger_car_NN, cmd_vel Twist, battery_voltage, LED" if c['code'] == 'C2' else
                ("diagnostico_microros.sh, ping 100 paq (Tabla 3), hz (Tabla 4)" if c['code'] == 'C3' else
                ("Simetría create/destroy, reconexión autónoma tras parada, Tabla 5" if c['code'] == 'C4' else
                "rqt_graph comentado, informe estructurado, Anexo A individual")))
            ]
            for c in CRITERIA
        ],
        [4.2, 1.3, 2.5, 5.0, 4.5],
    )

    # 5. Registro de Evidencias
    add_heading(doc, "5. Registro de Evidencias de la Práctica (E1–E8)")
    add_table(
        doc,
        ["Código", "Evidencia Requerida", "Localizador en Documento / Repositorio / Archivo"],
        [
            ["E1", "Captura de ejecución del agente micro-ROS con -v6 mostrando creación de cliente, sesión y participantes.", ""],
            ["E2", "Salida de terminal del Mini-Reto 1 (serial): medición con ros2 topic hz del heartbeat a 2 Hz y 10 Hz.", ""],
            ["E3", "Captura de ros2 node list (/burger_car_NN/base_controller) y ros2 topic info cmd_vel (Subscription: 1).", ""],
            ["E4", "Captura de teleoperación con ros2 topic pub --once y teleop_twist_keyboard, con activación física del LED.", ""],
            ["E5", "Salida de ejecución de diagnostico_microros.sh mostrando agente, puertos, ping al router y al ESP32.", ""],
            ["E6", "Resultados cuantitativos de ping (100 paquetes, -i 0.2) al ESP32: cerca y lejos del router (RTT y pérdida).", ""],
            ["E7", "Registro de estabilidad de telemetría con ros2 topic hz para periodos de 1000, 200 y 100 ms.", ""],
            ["E8", "Cronometraje y captura de reconexión autónoma del ESP32 tras detener y relanzar el agente UDP.", ""],
        ],
        [2, 11, 4.5],
    )

    # 6. Tablas Experimentales Llenas
    add_heading(doc, "6. Tablas de Registro Experimental (Diligenciadas por el Equipo)")
    add_note(doc, "Instrucciones: Registre los datos cuantitativos y parámetros técnicos obtenidos durante las fases del taller.")

    p = doc.add_paragraph("Tabla 1: Configuración de Dispositivos, Interfaces y Parámetros de Red")
    p.runs[0].bold = True
    add_table(
        doc,
        ["Parámetro", "Valor Configurado / Asignado", "Método de Verificación"],
        [
            ["Modelo de placa ESP32 utilizada", "", "Inspección física del chip"],
            ["Puerto serial USB de programación", "/dev/ttyUSB o /dev/ttyACM", "ls /dev/tty*"],
            ["Red Wi-Fi (SSID)", "ros2", "Configuración de firmware"],
            ["Dirección IP de la PC (Host Agente)", "", "hostname -I"],
            ["Puerto UDP del Agente micro-ROS", "8888", "Parámetro --port 8888"],
            ["Dirección IP asignada al ESP32", "", "Monitor Serie ('IP del ESP32: ...')"],
            ["Dirección MAC del ESP32", "", "Monitor Serie / router"],
            ["Namespace oficial del robot móvil", "/burger_car_", "Macro ROBOT_NAMESPACE"],
        ],
        [6.0, 5.5, 6.0],
    )

    p = doc.add_paragraph("Tabla 2: Métricas del Heartbeat Serial (Mini-Reto 1, Fase 1)")
    p.runs[0].bold = True
    add_table(
        doc,
        ["Configuración", "Periodo Configurado (ms)", "Frecuencia Media Medida (Hz)", "Mínimo (s)", "Máximo (s)", "Desviación Estándar (std dev)"],
        [
            ["Nominal 2 Hz", "500 ms", "Hz", "", "", ""],
            ["Nominal 10 Hz", "100 ms", "Hz", "", "", ""],
        ],
        [3.0, 3.5, 3.2, 2.5, 2.5, 2.8],
    )

    p = doc.add_paragraph("Tabla 3: Calidad de Enlace Inalámbrico hacia el ESP32 (ping -c 100 -i 0.2 <IP_ESP32>)")
    p.runs[0].bold = True
    add_table(
        doc,
        ["Ubicación del Robot / ESP32", "Distancia Aprox.", "RTT Mín (ms)", "RTT Prom (ms)", "RTT Máx (ms)", "Jitter mdev (ms)", "Pérdida (%)"],
        [
            ["Cerca del Router (< 2 m, línea de vista)", "< 2 m", "", "", "", "", "%"],
            ["Lejos del Router (> 8 m o con obstáculos)", "> 8 m", "", "", "", "", "%"],
        ],
        [5.0, 2.2, 2.0, 2.0, 2.0, 2.3, 2.0],
    )

    p = doc.add_paragraph("Tabla 4: Estabilidad de la Telemetría según Periodo de Publicación (ros2 topic hz -w 50)")
    p.runs[0].bold = True
    add_table(
        doc,
        ["Periodo (TELEMETRY_PERIOD_MS)", "Tasa Teórica", "Tasa Media Medida (Hz)", "Mínimo (s)", "Máximo (s)", "Desviación Estándar (std dev)"],
        [
            ["1000 ms", "1.0 Hz", "Hz", "", "", ""],
            ["200 ms", "5.0 Hz", "Hz", "", "", ""],
            ["100 ms", "10.0 Hz", "Hz", "", "", ""],
        ],
        [5.0, 2.5, 2.5, 2.5, 2.5, 2.5],
    )

    p = doc.add_paragraph("Tabla 5: Registro de Resiliencia y Máquina de Estados de Reconexión")
    p.runs[0].bold = True
    add_table(
        doc,
        ["Evento Experimental", "Acción Ejecutada", "Respuesta Firmware ESP32", "Estado en ros2 node list", "Tiempo Cronometrado"],
        [
            ["Pérdida del Agente", "Ctrl+C al agente en PC", "Falla ping, destroy_entities()", "Nodo desaparece del grafo", "—"],
            ["Relanzamiento del Agente", "Relanzar agente UDP :8888", "Detecta agente, create_entities()", "Nodo reaparece en el grafo", "s (sin reset)"],
            ["Pérdida de Señal Wi-Fi", "Apagar Wi-Fi / alejar", "Entra a reconexión Wi-Fi", "Nodo desaparece del grafo", "s tras reconectar"],
        ],
        [3.0, 3.5, 4.0, 4.0, 3.0],
    )

    # 7. Cuestionario Técnico
    add_heading(doc, "7. Cuestionario de Análisis Técnico (Respondido por el Equipo)")

    add_response_box(
        doc,
        "Pregunta 1: Ausencia del Agente en ros2 node list y Rol de XRCE-DDS\n"
        "¿Por qué al ejecutar ros2 node list aparece directamente /burger_car_NN/base_controller pero NO aparece ningún "
        "nodo denominado micro_ros_agent? Explique técnicamente cómo opera el Agente micro-ROS como puente entre "
        "Micro XRCE-DDS y DDS estándar.",
        lines=4,
    )

    add_response_box(
        doc,
        "Pregunta 2: Distinción Crítica entre Métricas del Monitor Web y Ping al ESP32\n"
        "Compare lo que mide el Dashboard del Monitor de Red web (http://localhost:8080) frente a lo que mide el "
        "comando ping <IP_ESP32>. ¿Por qué una latencia baja hacia el router no garantiza que la comunicación con "
        "el ESP32 esté libre de jitter o pérdidas? Explique los tramos involucrados.",
        lines=4,
    )

    add_response_box(
        doc,
        "Pregunta 3: Necesidad del Remapeo en teleop_twist_keyboard\n"
        "¿Por qué el nodo estándar teleop_twist_keyboard requiere obligatoriamente el remapeo cmd_vel:=/burger_car_NN/cmd_vel "
        "para operar el carrito? ¿Qué ocurriría en un laboratorio con múltiples carritos si todos operaran en el tópico raíz /cmd_vel?",
        lines=4,
    )

    add_response_box(
        doc,
        "Pregunta 4: Simetría de Entidades y Gestión de Memoria en C / C++ Embebido\n"
        "En el firmware, analice por qué es fundamental que create_entities() y destroy_entities() sean exactamente simétricas. "
        "Dado que el ESP32 cuenta con 520 KB de SRAM y no dispone de Garbage Collector, ¿qué problema ocurriría si no se liberan los recursos?",
        lines=4,
    )

    add_response_box(
        doc,
        "Pregunta 5: Impacto de la Tasa de Telemetría sobre Jitter y Uso de CPU\n"
        "A partir de los datos registrados en la Tabla 4, analice el comportamiento de la desviación estándar (std dev) "
        "al pasar de 1 Hz a 10 Hz. ¿Por qué en un sistema embebido con FreeRTOS y Wi-Fi no es conveniente elevar la frecuencia "
        "más allá de lo necesario?",
        lines=4,
    )

    add_response_box(
        doc,
        "Pregunta 6: Seguridad Funcional ante Pérdida de Comunicación\n"
        "En el Mini-Reto 3 se propone accionar motores reales. Si el robot móvil avanza a 0.5 m/s y se cae la red Wi-Fi o se "
        "apaga el agente, ¿qué le ocurriría a los motores si no se implementa una parada en destroy_entities()? Formule una política Fail-Safe.",
        lines=4,
    )

    # 8. Anexo A
    add_heading(doc, "8. Anexo A — Comprobación Individual de Desempeño")
    add_note(doc, "Propósito ABET: Verifica la autoría, comprensión técnica y contribución directa individual de cada estudiante.")

    for est_num in [1, 2]:
        p = doc.add_paragraph(f"Estudiante {est_num}: __________________________________________________ Código: ____________________")
        p.runs[0].bold = True
        add_table(
            doc,
            ["Pregunta Individual de Verificación", "Respuesta / Evidencia Directa del Estudiante"],
            [
                ["1. Rol y tareas técnicas desarrolladas: Describa sus responsabilidades específicas en la práctica (ej. compilación/ejecución del agente, conexión serial, Wi-Fi, teleoperación o diagnóstico).", ""],
                ["2. Arquitectura micro-ROS: Explique qué hace la capa rclc y el executor en el firmware del ESP32 frente al cliente XRCE-DDS.", ""],
                ["3. Máquina de estados de reconexión: Explique cómo el firmware detecta que se perdió el agente y qué pasos sigue para reconectarse sin pulsar reset.", ""],
                ["4. Análisis de una falla resuelta: Describa un problema técnico enfrentado (ej. puerto ocupado en Windows/WSL, AGENT_IP, firewall, compilación) y cómo lo solucionó.", ""],
                ["5. Autoría y reproducibilidad: Indique qué comandos, mediciones de tablas, capturas o secciones del firmware fueron elaborados directamente por usted.", ""],
            ],
            [7.5, 10.0],
        )
        p = doc.add_paragraph(f"Estado de comprobación individual Estudiante {est_num}: Verificada ☐     Insuficiente ☐     NA autorizado ☐     Firma: ________________________")
        p.paragraph_format.space_after = Pt(6)

    # 9. Rúbricas Docentes
    add_heading(doc, "9. Selección del Nivel Alcanzado por Criterio (Rúbricas Docentes)")
    for criterion in CRITERIA:
        add_heading(
            doc,
            f"{criterion['code']}. {criterion['name']} — {criterion['weight']}% — {criterion['so']}",
            level=2,
        )
        rows = []
        for level, descriptor in zip(LEVELS, criterion["descriptors"]):
            rows.append(["☐", f"{level[0]} — {level[1]}", descriptor])
        add_table(doc, ["Marque", "Nivel", "Evidencia observable"], rows, [1.5, 3, 12.5])
        doc.add_paragraph(f"Nivel {criterion['code']} marcado: ________     Valor Zubatronic (0–500): ________")

    # 10. Consolidado Final
    add_heading(doc, "10. Consolidado Final de Calificación Académica")
    add_table(
        doc,
        ["Criterio", "Peso Oficial (%)", "Nivel Marcado", "Valor Obtenido (0–500)", "Aporte Ponderado"],
        [
            [f"{c['code']}. {c['name']}", f"{c['weight']}%", "", "", ""] for c in CRITERIA
        ] + [["TOTAL CONSOLIDADO", "100%", "", "", "________ / 500"]],
        [6.5, 2.5, 2.2, 3.3, 3.0],
    )
    add_note(
        doc,
        "Nota Académica sobre 5,0 = Nota Consolidada sobre 500 ÷ 100. "
        "Aporte a Talleres del Segundo Corte: T₂ = Nota Académica sobre 5,0 (Componente T_C2).",
    )

    add_table(
        doc,
        ["Resultado Oficial de la Actividad", "Registro Oficial"],
        [
            ["Nota de Taller micro-ROS ESP32 sobre 500 puntos", "__________ / 500"],
            ["Nota Académica Oficial sobre 5,0", "__________ / 5,0"],
            ["Número de Criterios en Nivel N3 o superior (Umbral Individual)", "_____ / 5"],
            ["¿Cumple el Umbral Individual de Logro ABET (todos en N3+)?", "SÍ ☐         NO ☐"],
        ],
        [8.5, 9.0],
    )

    # 11. Consolidación ABET
    add_heading(doc, "11. Consolidación ABET y Cierre de Mejora Continua")
    add_table(
        doc,
        ["Student Outcome / Indicador Evaluado", "N Evaluables", "N en N3+", "% Logro", "Meta", "Hallazgo Docente"],
        [
            ["SO1 / Indicador 1.1 (Arquitectura Distribuida y Agente)", "", "", "", "70%", ""],
            ["SO2 / Indicador 2.2 (Diseño con Hardware Heterogéneo)", "", "", "", "70%", ""],
            ["SO6 / Indicador 6.1 y 6.4 (Experimentación y Diagnóstico)", "", "", "", "70%", ""],
            ["SO3 / Indicador 3.1 y 3.3 (Documentación y Trazabilidad)", "", "", "", "70%", ""],
            ["SO5 / Indicador 5.1 (Trabajo en Equipo y Roles)", "", "", "", "70%", ""],
        ],
        [5.0, 2.0, 2.0, 2.0, 1.8, 4.7],
    )

    add_table(
        doc,
        ["Campo de Cierre de Ciclo de Mejora Continua", "Registro de Gestión Pedagógica"],
        [
            ["Decisión derivada del hallazgo", ""],
            ["Acción de mejora continua pedagógica y técnica", ""],
            ["Responsable de la acción", "Ing. Henry Roncancio"],
            ["Fecha prevista de seguimiento", ""],
            ["Evidencia de seguimiento", ""],
            ["Resultado observado en el segundo ciclo", ""],
        ],
        [6.0, 11.5],
    )

    # 12. Cierre y Firmas
    add_heading(doc, "12. Cierre y Firmas de Conformidad")
    doc.add_paragraph("Docente Evaluador: Ing. Henry Roncancio           Fecha de Evaluación: ____________________")
    doc.add_paragraph("Firma del Docente Evaluador: __________________________________________________")
    doc.add_paragraph("Observaciones Finales y Retroalimentación Pedagógica:\n________________________________________________________________________________\n________________________________________________________________________________")

    add_page_number_footer(doc)
    doc.core_properties.title = "Instrumento ABET Taller micro-ROS ESP32 Robótica Móvil"
    doc.core_properties.subject = "ROBOT OPERATING SYSTEM - ROS · 2026-2"
    doc.core_properties.author = "Ing. Henry Roncancio"
    doc.core_properties.comments = "Versión 1.0. Documento único de entrega por equipos y evaluación ABET con tablas de métricas de red, cuestionario técnico y Anexo A."
    TARGET.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(TARGET))
    print(f"Instrumento ABET Word generado exitosamente en: {TARGET}")


if __name__ == "__main__":
    build_instrument()
