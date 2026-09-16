#!/usr/bin/env python3
"""
Constructor para INSTRUMENTO_ABET_PROYECTO_CORTE_2_CONTROL_POSICION.docx (Versión 1.1)
Genera el instrumento Word oficial de entrega estudiantil y evaluación ABET para el Proyecto
Integrador del Corte 2: Control de Posición en Lazo Cerrado con ROS 2 Actions (Turtlesim / micro-ROS).
Enfoque y ponderación reforzados en la arquitectura, ciclo de vida y contratos de Actions.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
TARGET = ROOT / "education" / "evidencias_abet" / "INSTRUMENTO_ABET_PROYECTO_CORTE_2_CONTROL_POSICION.docx"

FONT = "Arial"
BLUE = "1F4E79"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F2F2"

INDICATOR_21_23 = (
    "Indicador de desempeño 2.1 / 2.3: Diseña contratos de comunicación robustos integrando "
    "interfaces customizadas de acciones y arquitectura cliente-servidor en ROS 2."
)
INDICATOR_64 = (
    "Indicador de desempeño 6.4: Interpreta fallas y diagnósticos experimentales aplicando protocolos "
    "de diagnóstico por capas para aislar errores en hardware y software."
)
INDICATOR_11_12 = (
    "Indicador de desempeño 1.1 / 1.2: Analiza relaciones entre modelos matemáticos y orientación, "
    "calculando el error de rumbo y posición sin singularidades angulares."
)
INDICATOR_61_62 = (
    "Indicador de desempeño 6.1 / 6.2: Diseña y ejecuta pruebas de lazo cerrado, midiendo "
    "tiempos de convergencia y errores en estado estacionario."
)
INDICATOR_31_33_51 = (
    "Indicador de desempeño 3.1 - 3.3 / 5.1: Elabora documentación técnica reproducible, "
    "define roles técnicos y coordina la ejecución en equipo."
)

LEVELS = [
    ("N5", "475–500", "Excelente: dominio exhaustivo de ROS 2 Actions (contratos DDS, ciclo de vida formal, preempción en caliente y cancelación en seco), control cinemático sin singularidades y sustentación sobresaliente."),
    ("N4", "400–474", "Bueno: desempeño técnico correcto: la acción ejecuta metas, emite feedback a 10 Hz y cancela adecuadamente, con omisiones menores en preempción o en el análisis de jitter."),
    ("N3", "300–399", "Aceptable: demuestra el desempeño esencial con evidencia verificable: el servidor acepta la meta, calcula control y frena en tolerancia. Es el umbral individual de logro."),
    ("N2", "150–299", "Cumplimiento parcial: evidencia incompleta, acción implementada sin soporte de cancelación, inestabilidad en la trayectoria o sin introspección de tópicos internos."),
    ("N1", "0–149", "No cumple: implementación no funcional, sin arquitectura de acciones (uso de tópicos/servicios simples) o entrega fragmentaria. Sin evidencia obligatoria se registra 0."),
]

PROJECT_CRITERIA = [
    {
        "code": "C1",
        "name": "Arquitectura y contrato de Acción ROS 2",
        "weight": 30,
        "so": "Student Outcome SO2",
        "indicator": INDICATOR_21_23,
        "descriptors": [
            "Además de N4, implementa validación exhaustiva de precondiciones en goal_callback (rechazo por límites del lienzo [0.5, 10.5] o velocidades inválidas), introspecciona y documenta los 5 canales DDS generados por la acción, y explica la coordinación entre servicios y tópicos subyacentes.",
            "Compila exitosamente la interfaz GoToPose.action con los tipos requeridos, estructura callbacks asíncronos independientes (goal, cancel, execute) y valida la aceptación y rechazo de metas desde el ActionClient y por CLI.",
            "Define y compila el archivo .action, levanta el servidor y cliente de acción, y logra completar una meta básica en el lienzo.",
            "Presenta errores de compilación en rosidl o bloquea el hilo principal del nodo al omitir el bucle de ejecución asíncrono.",
            "No implementa la arquitectura de acciones (usa tópicos/servicios simples) o carece de evidencias funcionales.",
        ],
    },
    {
        "code": "C2",
        "name": "Ciclo de vida, máquina de estados y preempción",
        "weight": 25,
        "so": "Student Outcome SO6",
        "indicator": INDICATOR_64,
        "descriptors": [
            "Además de N4, implementa política de reemplazo de meta en caliente (Goal Preemption): ante la llegada de una nueva meta mientras ejecutaba, aborta ordenadamente la anterior y transiciona a la nueva sin detener el nodo, registrando transiciones formales en GoalStatus.",
            "Gestiona adecuadamente el ciclo de vida del GoalHandle, emite feedback periódico a 10 Hz exactos, detiene el robot ante cancel_goal_async() con parada en seco instantánea (v=0, w=0) y marca estado CANCELED.",
            "Publica feedback intermitente y frena el robot al alcanzar la tolerancia de distancia, aunque la orientación final queda fuera de tolerancia o no soporta cancelación.",
            "No soporta cancelación (ignora is_cancel_requested) o el robot sigue moviéndose indefinidamente tras la cancelación.",
            "Sin feedback periódico, sin verificación de tolerancias ni soporte de interrupción por software.",
        ],
    },
    {
        "code": "C3",
        "name": "Modelado cinemático en lazo cerrado",
        "weight": 20,
        "so": "Student Outcome SO1",
        "indicator": INDICATOR_11_12,
        "descriptors": [
            "Además de N4, desacopla formalmente las dos fases de control (rumbo y orientación final), demuestra analíticamente la estabilidad de Lyapunov del lazo, implementa atenuación suave con v*cos(alpha) en curvas cerradas y saturación cinemática sin discontinuidades.",
            "Calcula correctamente distancia euclidiana rho, error de rumbo alpha con wrap_to_pi y error angular theta_e, aplicando ganancias proporcionales sintonizadas que guían el robot a la meta suavemente.",
            "Implementa el cálculo cinemático de error cartesiano y orienta el robot hacia la meta, aunque presenta sobrepasos moderados o giros mayores a 180° por omisión de wrap_to_pi.",
            "La ley de control genera oscilaciones sostenidas, singularidades matemáticas cerca de la meta (rho -> 0) o velocidades fuera de los límites de seguridad.",
            "No formula el modelo cinemático del uniciclo ni la ley de control; comandos erráticos sin convergencia.",
        ],
    },
    {
        "code": "C4",
        "name": "Misión multi-waypoint y rosbag MCAP",
        "weight": 15,
        "so": "Student Outcome SO6",
        "indicator": INDICATOR_61_62,
        "descriptors": [
            "Además de N4, ejecuta una misión de patrullaje encadenando 4 waypoints secuenciales esperando el Result de cada uno antes de despachar el siguiente, graba rosbag MCAP de alta resolución y genera gráficas de convergencia y perfiles de velocidad.",
            "Ejecuta los escenarios experimentales básicos (rectilíneo, rotación pura, diagonal y cancelación), diligencia completamente las Tablas 2 y 3, y aporta capturas de trayectoria claras en turtlesim.",
            "Ejecuta al menos 2 escenarios experimentales, registra datos básicos en las tablas y aporta evidencia fotográfica del rastro de la tortuga.",
            "Presenta datos experimentales incompletos, sin dataset de rosbag verificable o con discrepancias notorias entre lo reportado y el comportamiento real.",
            "Sin pruebas experimentales documentadas, sin rosbag y sin métricas cuantitativas.",
        ],
    },
    {
        "code": "C5",
        "name": "Trazabilidad técnica, trabajo en equipo y Sim-to-Real",
        "weight": 10,
        "so": "Student Outcome SO3 / SO5",
        "indicator": INDICATOR_31_33_51,
        "descriptors": [
            "Además de N4, sustenta con fluidez técnica sobresaliente, demuestra commits Git equilibrados y trazables de todos los integrantes, y ejecuta con éxito el reto integrador Sim-to-Real remapeando el servidor hacia el robot físico con micro-ROS en ESP32.",
            "Entrega el informe técnico completo en formato oficial, responde con claridad a la sustentación técnica individual, mantiene el repositorio organizado y diligencia el Anexo A de comprobación.",
            "Presenta el informe técnico en el formato oficial, con Anexo A firmado y aportes individuales identificables.",
            "Documento desordenado, sin respuestas justificadas en el cuestionario o con participación pasiva de algún integrante en la sustentación.",
            "Falta el Anexo A individual, entrega incompleta o plagio/copia no autorizada.",
        ],
    },
]

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m, v in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)

def set_table_borders(table, color="D3D3D3"):
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement("w:tblBorders")
    for b in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{b}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)
        tblBorders.append(node)
    tblPr.append(tblBorders)

def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = FONT
    if level == 1:
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    elif level == 2:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
    else:
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    return p

def add_callout(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_BLUE)
    set_cell_margins(cell, 120, 120, 180, 180)
    tcPr = cell._tc.get_or_add_tcPr()
    tblBorders = OxmlElement("w:tcBorders")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "24")
    left.set(qn("w:color"), BLUE)
    tblBorders.append(left)
    for b in ("top", "bottom", "right"):
        node = OxmlElement(f"w:{b}")
        node.set(qn("w:val"), "none")
        tblBorders.append(node)
    tcPr.append(tblBorders)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(9.5)
    run.font.italic = True
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_table(doc, headers, data, col_widths=None):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    for c_idx, h in enumerate(headers):
        cell = table.cell(0, c_idx)
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell, 80, 80, 120, 120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(h)
        run.bold = True
        run.font.name = FONT
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for r_idx, row in enumerate(data):
        shading = LIGHT_GRAY if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx + 1, c_idx)
            set_cell_shading(cell, shading)
            set_cell_margins(cell, 70, 70, 100, 100)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(1.5)
            run = p.add_run(str(val))
            run.font.name = FONT
            run.font.size = Pt(9)
            if c_idx == 0:
                run.bold = True
    if col_widths:
        for r in table.rows:
            for c_idx, w in enumerate(col_widths):
                r.cells[c_idx].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table

def add_note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(f"Nota institucional: {text}")
    run.font.name = FONT
    run.font.size = Pt(8.5)
    run.font.italic = True
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

def build_document():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    # Título Principal
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("Instrumento de Evidencia y Evaluación ABET — Proyecto Integrador Corte 2")
    run_title.bold = True
    run_title.font.name = FONT
    run_title.font.size = Pt(16)
    run_title.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(8)
    run_sub = p_sub.add_run("Control de Posición en Lazo Cerrado con ROS 2 Actions (Turtlesim / micro-ROS) — Foco en Acciones")
    run_sub.font.name = FONT
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    add_callout(
        doc,
        "Documento Único de Entrega Estudiantil y Evaluación ABET. Asociado a education/proyectos_evaluables/PROYECTO_CORTE_2_CONTROL_POSICION_ACCIONES.md. "
        "Este formato evalúa el dominio de ROS 2 Actions: contratos DDS, máquina de estados formal, preempción en caliente, "
        "cancelación asíncrona, control cinemático polar y Anexo A individual de sustentación."
    )

    # 1. Identificación y Control
    add_heading(doc, "1. Identificación y Control")
    add_table(
        doc,
        ["Campo", "Registro Oficial"],
        [
            ["Programa Académico", "Ingeniería Mecatrónica"],
            ["Asignatura", "ROBOT OPERATING SYSTEM - ROS"],
            ["Periodo Académico", "2026-2"],
            ["Corte / Instrumento", "Segundo Corte / Entrega final y sustentación técnica (E_C2 — 30%)"],
            ["Actividad Evaluada", "Proyecto Integrador Corte 2 — Control de Posición con ROS 2 Actions (Turtlesim / micro-ROS)"],
            ["Número de Grupo / Subgrupo", ""],
            ["Estudiante 1 (Nombre y Código)", ""],
            ["Estudiante 2 (Nombre y Código)", ""],
            ["Estudiante 3 (si aplica)", ""],
            ["Plataforma de Validación", "Simulador Turtlesim ☐        micro-ROS ESP32 Físico (Bono) ☐"],
            ["Nombre del Archivo de Entrega", "C2_E02_G<grupo>_<codigo1>_<codigo2>_v1.docx"],
            ["Fecha de Realización / Sustentación", ""],
            ["Fecha de Entrega del Documento", ""],
            ["Docente Evaluador", "Ing. Henry Roncancio"],
            ["Versión del Instrumento", "Versión 1.1 — Enfoque Profundo en Acciones (2026-2)"],
            ["Unidad de Análisis / Captura", "Equipo colaborativo con comprobación individual (Anexo A)"],
        ],
        [5, 12],
    )

    # 2. Parámetros de assessment
    add_heading(doc, "2. Parámetros de Assessment")
    add_table(
        doc,
        ["Parámetro", "Regla Institucional y Metodológica Adoptada"],
        [
            ["Población o cohorte", "Censo completo de estudiantes matriculados que presentan el proyecto integrador del Corte 2 en 2026-2."],
            ["Momento de medición", "Segundo corte (Semana 13), tras la implementación de la interfaz de acción, servidor cinemático y pruebas."],
            ["Evaluador", "Docente titular de la asignatura ROBOT OPERATING SYSTEM - ROS."],
            ["Umbral individual de logro", "Nivel N3 o superior (puntaje mínimo de 300 sobre 500) en cada indicador evaluado."],
            ["Meta de cohorte", "Al menos el 70% de los estudiantes evaluables debe alcanzar el nivel N3 o superior en cada indicador."],
            ["Regla de muestreo", "Censo 100%: no se utiliza muestreo; se evalúa la totalidad de los estudiantes y equipos exigibles."],
            ["Evidencia faltante", "Entrega exigible sin evidencia obligatoria verificable: N1 con valor 0. Retiro oficial: NA / no evaluado."],
            ["Regla de trabajo colaborativo", "Informe único por equipo con tablas llenas más Anexo A individual obligatorio."],
        ],
        [5, 12],
    )

    # 3. Niveles de desempeño
    add_heading(doc, "3. Niveles de Desempeño para Zubatronic/SGDE")
    add_table(
        doc,
        ["Nivel", "Intervalo Zubatronic", "Interpretación y Criterio de Logro"],
        [[lvl, r, desc] for lvl, r, desc in LEVELS],
        [2.5, 3.5, 11],
    )

    # 4. Alineación de Criterios (Rebalanceada con Foco en Actions)
    add_heading(doc, "4. Alineación de Criterios, RAE y Student Outcomes (Ponderación Reforzada en Actions)")
    add_table(
        doc,
        ["Criterio", "Peso", "SO", "Indicador del Programa", "Evidencia Directa Obligatoria"],
        [
            ["C1. Arquitectura y contrato de Acción", "30%", "SO2", INDICATOR_21_23, "Definición GoToPose.action compilada, 5 canales DDS identificados y validación en goal_callback."],
            ["C2. Ciclo de vida y preempción", "25%", "SO6", INDICATOR_64, "Máquina de estados GoalStatus, emisión a 10 Hz, parada ante cancelación y política de reemplazo en caliente."],
            ["C3. Modelado cinemático polar", "20%", "SO1", INDICATOR_11_12, "Ley de control rho/alpha/theta_e, wrap_to_pi, saturaciones, y convergencia en dos fases."],
            ["C4. Misión multi-waypoint y rosbag", "15%", "SO6", INDICATOR_61_62, "Misión de patrullaje encadenada, dataset MCAP válido, curvas de error vs tiempo y velocidades."],
            ["C5. Trazabilidad técnica y trabajo en equipo", "10%", "SO3/5", INDICATOR_31_33_51, "Informe técnico en formato oficial, repositorio Git con commits trazables, Anexo A y sustentación."],
            ["TOTAL", "100%", "", "", "500 Puntos Máximos (Nota = Puntos / 100)"],
        ],
        [3.8, 1.4, 1.8, 5.0, 5.0],
    )

    # 5. Registro de evidencias
    add_heading(doc, "5. Registro de Evidencias Obligatorias (E1–E8)")
    add_table(
        doc,
        ["Código", "Evidencia Requerida", "Localizador en el Documento / Repositorio"],
        [
            ["E1", "Captura de ros2 interface show GoToPose.action mostrando Goal, Result y Feedback.", ""],
            ["E2", "Salida de terminal de la introspección de los 5 canales DDS (ros2 topic/service list | grep go_to_pose).", ""],
            ["E3", "Captura del servidor mostrando aceptación de metas válidas y rechazo de metas fuera de límites.", ""],
            ["E4", "Captura del cliente mostrando recepción periódica de Feedback a 10 Hz con distancia y error angular.", ""],
            ["E5", "Evidencia de la prueba de Preempción en caliente (Goal Replacement fluido sin detener el nodo).", ""],
            ["E6", "Evidencia de la prueba de Cancelación voluntaria demostrando parada instantánea y STATUS_CANCELED.", ""],
            ["E7", "Captura de pantalla de turtlesim con el rastro de la Misión Multi-Waypoint (4 waypoints encadenados).", ""],
            ["E8", "Salida de rosbag info del dataset MCAP y gráficas de trayectoria y velocidad generadas.", ""],
        ],
        [1.8, 9.2, 6.0],
    )

    # 6. Tablas experimentales
    add_heading(doc, "6. Tablas de Registro Experimental (Diligenciadas por el Equipo)")

    add_heading(doc, "Tabla 1: Introspección de los Canales DDS de la Acción /go_to_pose", level=2)
    add_table(
        doc,
        ["Canal Generado por ROS 2", "Tipo de Primitiva", "Tipo de Mensaje / Servicio", "Función en el Ciclo de Vida"],
        [
            ["/go_to_pose/_action/send_goal", "Servicio", "<paquete>/action/GoToPose_SendGoal", "Validación y aceptación/rechazo de la meta"],
            ["/go_to_pose/_action/cancel_goal", "Servicio", "action_msgs/srv/CancelGoal", "Solicitud de interrupción en caliente"],
            ["/go_to_pose/_action/get_result", "Servicio", "<paquete>/action/GoToPose_GetResult", "Espera asíncrona del desenlace final"],
            ["/go_to_pose/_action/feedback", "Tópico", "<paquete>/action/GoToPose_FeedbackMessage", "Publicación periódica del avance (10 Hz)"],
            ["/go_to_pose/_action/status", "Tópico", "action_msgs/msg/GoalStatusArray", "Array con estados de los GoalHandle activos"],
        ],
        [5.5, 2.5, 4.5, 4.5],
    )

    add_heading(doc, "Tabla 2: Resultados Cuantitativos de los Escenarios Experimentales", level=2)
    add_table(
        doc,
        ["Escenario", "Pose Inicial [x, y, th]", "Pose Meta [x, y, th]", "Error rho [m]", "Error th [rad]", "Tiempo [s]", "Estado Final Reportado"],
        [
            ["E1: Rectilíneo", "[5.54, 5.54, 0.00]", "[9.50, 5.54, 0.00]", "", "", "", "STATUS_SUCCEEDED ☐"],
            ["E2: Rotación", "[5.54, 5.54, 0.00]", "[5.54, 5.54, 3.14]", "", "", "", "STATUS_SUCCEEDED ☐"],
            ["E3: Patrulla W1-W4", "Misión 4 Waypoints", "Ruta Poligonal", "", "", "", "Todos SUCCEEDED ☐"],
            ["E4: Cancelación", "[2.00, 2.00, 0.00]", "[10.0, 10.0, 0.00]", "N/A", "N/A", "", "STATUS_CANCELED ☐"],
            ["E5: Preempción", "Meta 1 -> Meta 2 en vuelo", "M1 -> M2", "", "", "", "M1 ABORT, M2 SUCC ☐"],
        ],
        [3.0, 3.0, 3.0, 2.0, 2.0, 1.8, 2.2],
    )

    add_heading(doc, "Tabla 3: Verificación de Robustez y Manejo de Estados de la Acción", level=2)
    add_table(
        doc,
        ["Prueba de Robustez", "Entrada Aplicada", "Comportamiento Esperado", "Comportamiento Real", "¿Aprobado?"],
        [
            ["Rechazo fuera de límites", "x = 12.0, y = 5.0", "Rechazo en goal_callback (GoalResponse.REJECT)", "", "SÍ ☐   NO ☐"],
            ["Parada en cancelación", "cancel_goal() en vuelo", "Frenado en seco v=0, w=0, estado CANCELED", "", "SÍ ☐   NO ☐"],
            ["Preempción concurrente", "Meta 2 recibida a t=2s", "Aborto ordenado de Meta 1, transición a Meta 2", "", "SÍ ☐   NO ☐"],
            ["Cadencia de Feedback", "ros2 topic hz", "10 Hz +/- 1 Hz en /go_to_pose/_action/feedback", "", "SÍ ☐   NO ☐"],
        ],
        [3.5, 3.2, 4.5, 4.3, 1.5],
    )

    # 7. Cuestionario de análisis
    add_heading(doc, "7. Cuestionario de Análisis Técnico y Justificación de Ingeniería")
    questions = [
        "1. La Tríada de Comunicación: Compare conceptual y arquitectónicamente Tópicos, Servicios y Acciones en ROS 2. ¿Bajo qué criterios de ingeniería se decide que una tarea robótica debe implementarse como una Acción en lugar de un Servicio o un Tópico?",
        "2. Los 5 Canales DDS: Explique en detalle qué sucede en la capa de red cuando un cliente envía una meta con send_goal_async(). ¿Por qué la Acción genera 3 servicios y 2 tópicos en lugar de uno solo?",
        "3. Máquina de Estados y Preempción: Describa los estados del GoalStatus de ROS 2. Explique cómo su nodo servidor gestiona la llegada de una nueva meta mientras otra se encuentra en estado STATUS_EXECUTING y por qué esta política es vital en robots de navegación autónoma (Nav2).",
        "4. Manejo de Cancelación Asíncrona: ¿Cuál es la diferencia entre que el middleware acepte la cancelación en cancel_callback y que el hilo de ejecución confirme la detención física mediante goal_handle.canceled()? ¿Por qué nunca debe matarse el proceso bruscamente?",
        "5. Portabilidad Sim-to-Real con micro-ROS: Justifique técnicamente cómo la independencia de transporte de ROS 2 permite que el mismo servidor de acciones comande al robot móvil físico con ESP32 simplemente remapeando el tópico cmd_vel.",
    ]
    for q in questions:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(q)
        r.font.name = FONT
        r.font.size = Pt(9.5)
        r.bold = True
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.space_before = Pt(0)
        p_ans.paragraph_format.space_after = Pt(8)
        r_ans = p_ans.add_run("[Espacio para respuesta técnica, ecuaciones y justificación del equipo]\n___________________________________________________________________________________________________")
        r_ans.font.name = FONT
        r_ans.font.size = Pt(9)
        r_ans.font.italic = True
        r_ans.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

    # 8. Anexo A
    add_heading(doc, "8. Anexo A: Comprobación Individual del Logro ABET (Obligatorio)")
    for est_num in [1, 2]:
        add_heading(doc, f"Estudiante {est_num}: _____________________________________________ Código: ______________", level=2)
        add_table(
            doc,
            ["Campo de Verificación Individual", "Registro del Estudiante / Evaluación Docente"],
            [
                ["Rol técnico asumido en el proyecto", "Arquitectura de acción y callbacks ☐    Control cinemático y preempción ☐    Cliente interactivo y waypoints ☐    Rosbag y métricas ☐"],
                ["Contribución concreta al código (archivos y funciones)", ""],
                ["Pregunta de sustentación individual formulada por el docente", ""],
                ["Respuesta y justificación técnica del estudiante", ""],
                ["Nivel de logro individual asignado por el docente", "N1 ☐    N2 ☐    N3 (Umbral) ☐    N4 ☐    N5 ☐"],
            ],
            [5.5, 11.5],
        )

    # 9. Rúbricas analíticas docentes
    add_heading(doc, "9. Rúbricas Analíticas de Evaluación Docente (Ponderación Reforzada en Actions)")
    for crit in PROJECT_CRITERIA:
        add_heading(doc, f"{crit['code']}. {crit['name']} ({crit['weight']}%) — {crit['so']}", level=2)
        p_ind = doc.add_paragraph()
        p_ind.paragraph_format.space_after = Pt(2)
        r_ind = p_ind.add_run(crit["indicator"])
        r_ind.font.name = FONT
        r_ind.font.size = Pt(8.5)
        r_ind.font.italic = True
        rubric_data = [
            [lvl, r, desc]
            for (lvl, r, _), desc in zip(LEVELS, crit["descriptors"])
        ]
        add_table(doc, ["Nivel", "Intervalo", "Descriptor de Desempeño Observable"], rubric_data, [2.0, 2.5, 12.5])

    # 10. Consolidación de calificación
    add_heading(doc, "10. Consolidación de Calificación Docente (Escala Zubatronic 0–500)")
    add_table(
        doc,
        ["Criterio Evaluado", "Peso Oficial (%)", "Nivel Marcado", "Valor Obtenido (0–500)", "Aporte Ponderado"],
        [
            [f"{c['code']}. {c['name']}", f"{c['weight']}%", "", "", ""] for c in PROJECT_CRITERIA
        ] + [["TOTAL CONSOLIDADO", "100%", "", "", "________ / 500"]],
        [6.5, 2.5, 2.2, 3.3, 3.0],
    )
    add_note(
        doc,
        "Nota Académica sobre 5,0 = Nota Consolidada sobre 500 ÷ 100. "
        "Aporte a la Nota del Corte 2: E₂ = Nota Académica × 0,30 (30% de la nota de C2).",
    )

    add_table(
        doc,
        ["Resultado Oficial de la Actividad", "Registro Oficial"],
        [
            ["Nota del Proyecto Integrador Corte 2 sobre 500 puntos", "__________ / 500"],
            ["Nota Académica Oficial sobre 5,0", "__________ / 5,0"],
            ["Número de Criterios en Nivel N3 o superior (Umbral Individual)", "_____ / 5"],
            ["¿Cumple el Umbral Individual de Logro ABET (todos en N3+)?", "SÍ ☐         NO ☐"],
        ],
        [7.0, 10.0],
    )

    TARGET.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(TARGET))
    print(f"Instrumento ABET Word generado exitosamente en: {TARGET}")

if __name__ == "__main__":
    build_document()
