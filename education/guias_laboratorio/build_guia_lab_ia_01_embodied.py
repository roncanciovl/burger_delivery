#!/usr/bin/env python3
"""
Constructor de los Word de la Guía de Laboratorio IA 01 (Embodied AI con Gemini Robotics-ER 2).

La fuente única es GUIA_LAB_IA_01_EMBODIED_AI_GEMINI_ER2_KINOVA.md. Este script la convierte
al formato institucional GL-AA-F-1 (templates/Formato_Guias_de_Laboratorio.docx):

    GUIA_LAB_IA_01_EMBODIED_AI_GEMINI_ER2_KINOVA.docx

La entrega de los grupos es el instrumento ABET, que se genera aparte con
education/evidencias_abet/build_instrumento_lab_ia_01.py (reutiliza bloques() y anchos()).

Reutiliza los estilos del constructor de la Guía 02 para que todas las guías se vean igual.
Ejecución:  python3 education/guias_laboratorio/build_guia_lab_ia_01_embodied.py
"""

from __future__ import annotations

import re
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

AQUI = Path(__file__).resolve().parent
sys.path.insert(0, str(AQUI))
import build_guia_lab_02_camara_kinova as estilo  # noqa: E402

FUENTE = AQUI / "GUIA_LAB_IA_01_EMBODIED_AI_GEMINI_ER2_KINOVA.md"
GUIA = AQUI / "GUIA_LAB_IA_01_EMBODIED_AI_GEMINI_ER2_KINOVA.docx"

ASIGNATURA = "INTELIGENCIA ARTIFICIAL"
TITULO = ("Práctica IA 01. Embodied AI: razonamiento espacial con Gemini Robotics-ER 2 "
          "y localización 3D con la cámara RGB-D del Kinova Gen3")
FECHA_EMISION = "2026/09/28"
ETIQUETAS = {"WARNING": "ADVERTENCIA", "IMPORTANT": "IMPORTANTE", "NOTE": "NOTA", "TIP": "CONSEJO"}
SALTO_ANTES = ("MATERIALES", "PROCEDIMIENTO", "RESULTADOS")


# ---------------------------------------------------------------------------
# Markdown -> bloques
# ---------------------------------------------------------------------------
def latex_a_texto(texto: str) -> str:
    """Aproximación legible de las fórmulas LaTeX de la guía para Word."""
    t = texto
    t = re.sub(r"\\begin\{bmatrix\}(.*?)\\end\{bmatrix\}",
               lambda m: "[" + ", ".join(p.strip() for p in m.group(1).split("\\\\")) + "]", t)

    def fraccion(m):
        num, den = (x if re.fullmatch(r"[\w.]+", x) else f"({x})" for x in (m.group(1), m.group(2)))
        return f"{num}/{den}"

    t = t.replace(r"\tfrac", r"\frac")
    for _ in range(3):
        t = re.sub(r"\\frac\{([^{}]*(?:\{[^{}]*\}[^{}]*)*)\}\{([^{}]*)\}", fraccion, t)
    t = re.sub(r"\\operatorname\{([^{}]*)\}", r"\1", t)
    t = re.sub(r"\\sqrt\{([^{}]*)\}", r"√(\1)", t)
    reemplazos = {r"\qquad": "     ", r"\,": " ", r"\mid": "|", r"\sim": "~", r"\approx": "≈",
                  r"\cdot": "·", "^{-1}": "⁻¹", r"\{": "{", r"\}": "}",
                  r"\big": "", r"\lvert": "|", r"\rvert": "|", r"\le": "≤", r"\ge": "≥",
                  r"\pm": "±", r"\mp": "∓", r"\theta": "θ", r"\alpha": "α", r"\rho": "ρ",
                  r"\omega": "ω", r"\Delta": "Δ", r"\pi": "π", r"\sin": "sin", r"\cos": "cos",
                  r"\max": "max", "^2": "²", "\\ ": " "}
    for a, b in reemplazos.items():
        t = t.replace(a, b)
    t = re.sub(r"_\{([^{}]*)\}", r"_\1", t)
    return t.replace("$", "")


def texto_plano(texto: str) -> str:
    texto = re.sub(r"<br\s*/?>", "\n", texto)
    return latex_a_texto(texto)


def celdas(linea: str) -> list[str]:
    # «\|» es una barra literal dentro de una celda (p. ej. «cmd \| grep x»), no un separador.
    partes = re.split(r"(?<!\\)\|", linea.strip().strip("|"))
    return [texto_plano(c.strip().replace("\\|", "|")) for c in partes]


def bloques(md: str):
    """Genera (tipo, datos) para cada bloque del Markdown."""
    lineas = md.splitlines()
    i = 0
    while i < len(lineas):
        linea = lineas[i]
        s = linea.strip()
        if not s or s == "---":
            i += 1
        elif s.startswith("```"):
            j = i + 1
            while not lineas[j].strip().startswith("```"):
                j += 1
            yield "codigo", lineas[i + 1:j]       # el código va literal: «$HOME» no es LaTeX
            i = j + 1
        elif s.startswith("$$"):
            j = i if s.endswith("$$") and len(s) > 2 else i + 1
            while not lineas[j].strip().endswith("$$"):
                j += 1
            yield "formula", latex_a_texto(" ".join(x.strip() for x in lineas[i:j + 1]))
            i = j + 1
        elif s.startswith("#"):
            nivel = len(s) - len(s.lstrip("#"))
            yield f"h{nivel}", s.lstrip("#").strip()
            i += 1
        elif s.startswith(">"):
            contenido = []
            while i < len(lineas) and lineas[i].strip().startswith(">"):
                contenido.append(re.sub(r"^\s*>\s?", "", lineas[i]))
                i += 1
            m = re.match(r"\[!(\w+)\]", contenido[0].strip())
            etiqueta = ETIQUETAS.get(m.group(1), "NOTA") if m else "NOTA"
            cuerpo = contenido[1:] if m else contenido
            lineas_aviso = [texto_plano(c) for c in cuerpo if c.strip() and not c.strip().startswith("```")]
            yield "aviso", (etiqueta, "\n".join(lineas_aviso))
        elif s.startswith("|"):
            filas = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                if not re.match(r"^\|[\s:|-]+\|$", lineas[i].strip()):
                    filas.append(celdas(lineas[i]))
                i += 1
            yield "tabla", filas
        elif re.match(r"^\s*(\d+\.|-)\s", linea):
            items = []
            while i < len(lineas) and (re.match(r"^\s*(\d+\.|-)\s", lineas[i])
                                       or (lineas[i].startswith("   ") and lineas[i].strip()
                                           and not lineas[i].strip().startswith("```"))):
                m = re.match(r"^(\s*)(\d+\.|-)\s+(.*)", lineas[i])
                if m:
                    num = int(m.group(2)[:-1]) if m.group(2)[:-1].isdigit() else None
                    items.append([len(m.group(1)) // 3, num, m.group(3)])
                else:
                    items[-1][2] += " " + lineas[i].strip()
                i += 1
            yield "lista", [(n, num, texto_plano(t)) for n, num, t in items]
        else:
            parrafo = [s]
            i += 1
            while (i < len(lineas) and lineas[i].strip()
                   and not re.match(r"^\s*([#>|`$]|\d+\.\s|-\s)", lineas[i]) and lineas[i].strip() != "---"):
                parrafo.append(lineas[i].strip())
                i += 1
            yield "parrafo", texto_plano(" ".join(parrafo))


def secciones(md: str) -> dict[str, list]:
    """Agrupar los bloques por sección de nivel 2 (clave: título sin número)."""
    salida, actual = {}, "_inicio"
    for tipo, datos in bloques(md):
        if tipo == "h2":
            actual = re.sub(r"^\d+\.\s*", "", datos).upper()
            salida[actual] = []
        else:
            salida.setdefault(actual, []).append((tipo, datos))
    return salida


# ---------------------------------------------------------------------------
# Escritura en Word
# ---------------------------------------------------------------------------
def anchos(filas: list[list[str]], total_cm: float = 17.0) -> list[float]:
    """Repartir el ancho de página según el texto más largo de cada columna."""
    columnas = max(len(f) for f in filas)
    peso = [max(len(f[j]) if j < len(f) else 0 for f in filas) ** 0.5 + 2 for j in range(columnas)]
    # Una columna vacía es espacio para escribir a mano: darle ancho, no lo mínimo.
    vacias = [j for j in range(columnas) if not any(j < len(f) and f[j].strip() for f in filas[1:])]
    if len(filas) > 1 and len(vacias) < columnas:
        for j in vacias:
            peso[j] = max(peso[j], 0.8 * max(peso))
    return [round(total_cm * w / sum(peso), 2) for w in peso]


def fijar_anchos(tabla, anchos_cm: list[float]):
    """Fijar la cuadrícula de columnas además del ancho de celda: LibreOffice y Word
    ignoran el ancho de celda si la tabla queda en ajuste automático."""
    tabla.autofit = False
    for j, ancho in enumerate(anchos_cm):
        tabla.columns[j].width = Cm(ancho)
        for celda in tabla.columns[j].cells:
            celda.width = Cm(ancho)


def encabezado_seccion(doc, titulo: str, nueva_pagina: bool = False):
    """Como add_section_heading, pero el salto va en el propio título: así no quedan
    páginas en blanco cuando el contenido anterior llena la página justa."""
    p = estilo.add_section_heading(doc, titulo)
    p.paragraph_format.page_break_before = nueva_pagina
    return p


def mantener_junta(tabla):
    """Evitar que una tabla pequeña (firmas, aprobación) se parta entre dos páginas."""
    for fila in tabla.rows[:-1]:
        estilo.prevent_row_split(fila)
        for celda in fila.cells:
            for p in celda.paragraphs:
                p.paragraph_format.keep_with_next = True


def escribir(doc, contenido, modelo_tabla):
    for tipo, datos in contenido:
        if tipo in ("h3", "h4"):
            estilo.add_subheading(doc, re.sub(r"^\d+(\.\d+)*\.\s*", "", datos))
        elif tipo == "parrafo":
            estilo.add_body(doc, datos)
        elif tipo == "formula":
            estilo.add_body(doc, datos, align=WD_ALIGN_PARAGRAPH.CENTER, italic=True)
        elif tipo == "codigo":
            estilo.add_code_box(doc, datos)
        elif tipo == "aviso":
            estilo.add_callout(doc, *datos)
        elif tipo == "lista":
            for nivel, num, texto in datos:
                estilo.add_list_item(doc, texto, num, level=nivel)
        elif tipo == "tabla":
            fijar_anchos(estilo.add_data_table(doc, datos, modelo_tabla, font_size=7.5 if len(datos[0]) > 4 else 8),
                         anchos(datos))


def pagina_inicial(doc, cuerpo, sect_pr, modelos, docente: str, tipo: str | None = None):
    cover_header, cover_id, cover_sign = modelos
    cuerpo.insert(cuerpo.index(sect_pr), deepcopy(cover_header))
    cover = doc.tables[-1]
    if tipo:  # la guía conserva el rótulo institucional; el informe lo cambia
        estilo.replace_text_preserving_cell(cover.cell(0, 0), tipo, bold=True, size=18,
                                            align=WD_ALIGN_PARAGRAPH.CENTER)
    estilo.replace_text_preserving_cell(cover.cell(0, 1), f"Fecha Emisión:\n{FECHA_EMISION}", bold=True,
                                        size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    estilo.replace_text_preserving_cell(cover.cell(1, 1), "Revisión No.:\n1", bold=True, size=8,
                                        align=WD_ALIGN_PARAGRAPH.CENTER)
    celda = cover.cell(1, 2)
    celda.text = ""
    p = celda.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for texto, campo in (("Página ", "PAGE"), (" de ", "NUMPAGES")):
        r = p.add_run(texto)
        r.font.name, r.font.size, r.bold = estilo.FONT, Pt(8), True
        estilo.add_field(r, campo)
    estilo.add_body(doc, "", size=4)

    cuerpo.insert(cuerpo.index(sect_pr), deepcopy(cover_id))
    ident = doc.tables[-1]
    estilo.replace_text_preserving_cell(ident.cell(0, 0), f"Laboratorio de: {ASIGNATURA}", bold=True, size=9)
    estilo.replace_text_preserving_cell(ident.cell(1, 0), f"Título de Laboratorio: {TITULO}", bold=True, size=9)
    estilo.add_body(doc, "", size=6)

    if cover_sign is not None:
        cuerpo.insert(cuerpo.index(sect_pr), deepcopy(cover_sign))
        firmas = doc.tables[-1]
        for j, texto in enumerate((f"Elaborado por:\n\n{docente}",
                                   "Revisado por:\n\nDirector de Programa\nIngeniería Mecatrónica",
                                   "Aprobado por:\n\nDecano(a)\nFacultad de Ingeniería")):
            estilo.replace_text_preserving_cell(firmas.cell(0, j), texto, size=8, align=WD_ALIGN_PARAGRAPH.CENTER)


def cargar_modelos():
    doc = Document(estilo.TEMPLATE)
    t = doc.tables
    modelos = {"cover": (deepcopy(t[0]._tbl), deepcopy(t[1]._tbl), deepcopy(t[2]._tbl)),
               "cambios": deepcopy(t[3]._tbl), "aprobacion": deepcopy(t[6]._tbl)}
    material = t[4]
    estilo.clear_body(doc)
    return doc, modelos, material


def construir_guia(secs: dict, docente: str):
    doc, modelos, material = cargar_modelos()
    cuerpo = doc._element.body
    sect_pr = cuerpo.sectPr
    pagina_inicial(doc, cuerpo, sect_pr, modelos["cover"], docente)

    estilo.add_page_break(doc)
    estilo.add_section_heading(doc, "Control de Cambios de la Guía de práctica")
    cuerpo.insert(cuerpo.index(sect_pr), deepcopy(modelos["cambios"]))
    cambios = doc.tables[-1]
    filas = next(d for tipo, d in secs["CONTROL DE CAMBIOS"] if tipo == "tabla")[1:]
    for row in cambios.rows[1:]:
        for cell in row.cells:
            estilo.replace_text_preserving_cell(cell, "", size=8)
    for i, valores in enumerate(filas, start=1):
        for j, valor in enumerate(valores[:3]):
            estilo.replace_text_preserving_cell(cambios.cell(i, j), valor, size=8)
    estilo.set_repeat_header(cambios.rows[0])

    estilo.add_page_break(doc)
    encabezado = next(d for tipo, d in secs["_inicio"] if tipo == "tabla")
    datos = dict(zip(encabezado[0], encabezado[1]))
    for etiqueta, clave in (("FACULTAD O UNIDAD ACADÉMICA:", "FACULTAD"), ("PROGRAMA:", "PROGRAMA"),
                            ("ASIGNATURA:", "ASIGNATURA"), ("SEMESTRE:", "SEMESTRE")):
        p = estilo.add_body(doc, "", keep=True)
        for texto, negrita in ((etiqueta + " ", True), (datos.get(clave, ""), False)):
            r = p.add_run(texto)
            r.font.name, r.font.size, r.bold = estilo.FONT, Pt(9), negrita
    # Avisos previos a la sección 1 (p. ej. sincronizar con la base del curso).
    escribir(doc, [b for b in secs["_inicio"] if b[0] == "aviso"], material)

    for titulo, contenido in secs.items():
        if titulo in ("_inicio", "CONTROL DE CAMBIOS") or titulo.startswith("APROBACIÓN"):
            continue
        encabezado_seccion(doc, titulo, nueva_pagina=titulo.startswith(SALTO_ANTES))
        escribir(doc, contenido, material)

    cuerpo.insert(cuerpo.index(sect_pr), deepcopy(modelos["aprobacion"]))
    aprobacion = doc.tables[-1]
    mantener_junta(aprobacion)
    for j, texto in enumerate((docente, "Director de Programa\nIngeniería Mecatrónica",
                               "Decano(a)\nFacultad de Ingeniería")):
        estilo.replace_text_preserving_cell(aprobacion.cell(2, j), texto, size=8, align=WD_ALIGN_PARAGRAPH.CENTER)

    estilo.configure_footer(doc)
    cp = doc.core_properties
    cp.title = "Guía de Laboratorio IA 01 – Embodied AI con Gemini Robotics-ER 2 y Kinova Gen3"
    cp.subject, cp.author = ASIGNATURA, "Universidad Militar Nueva Granada"
    cp.comments = f"Generado desde {FUENTE.name} con {Path(__file__).name}."
    doc.save(GUIA)
    print(f"Guía Word: {GUIA}")


def main():
    secs = secciones(FUENTE.read_text(encoding="utf-8"))
    docente = "Ing. Henry Roncancio\nDocente Asignatura Inteligencia Artificial"
    construir_guia(secs, docente)


if __name__ == "__main__":
    main()
